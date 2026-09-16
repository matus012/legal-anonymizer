"""Build the v1.1 DEMO document and drive it through the GUI's own pipeline.

    .\\.venv\\Scripts\\python.exe -m tools.make_demo

Dev-only. Produces a realistic synthetic Slovak kúpna zmluva with EVERY v1.1 type seeded,
in both .docx and .pdf, then runs the exact functions the GUI wizard calls — ``scan_file`` →
``build_decisions`` → ``export_file`` — and prints what a reviewer would see.

Why the model layer rather than clicking the Qt wizard: ``gui/app.py`` is a thin Qt shell over
``gui/model.py``. The scan page calls ``scan_file``, the review table is built from the
``ReviewRow``s it returns, and the export button calls ``export_file`` with the decisions the
table produced. Driving the model directly exercises the same code on the same document and,
unlike a screenshot, leaves an artifact that can be re-run and diffed. The Qt shell itself was
verified live by the owner in Phase 6 and is unchanged by this sprint except for the review
table's new checksum column.

The names, numbers and addresses below are INVENTED. No real client data is used anywhere in
this project (context.md §7) — including in the demo.
"""
from __future__ import annotations

import os
import shutil
from pathlib import Path

import fitz
from docx import Document

from corpus.pdf_builder import _find_font
from gui.model import build_decisions, export_file, out_path_for, scan_file

DEMO_DIR = Path("demo")

# The parties the lawyer would type into the "known entities" box (context.md §4.3).
KNOWN_ENTITIES = ["Ján Novák", "Mária Kováčová"]

# Every paragraph of the demo contract. Each line is placed verbatim; the types it is meant to
# exercise are named in the comment so a reviewer can check coverage against the report.
PARAGRAPHS = [
    "Kúpna zmluva",
    "uzavretá podľa § 588 a nasl. Občianskeho zákonníka",
    "",
    "Článok I — Zmluvné strany",
    "Predávajúci: JUDr. Ján Novák",                              # MENO (title + role anchor)
    "Rodné číslo: 850315/0018",                                  # RODNE_CISLO (valid)
    "Dátum narodenia: 15.3.1985",                                # DATUM (DOB-anchored)
    "Trvale bytom: Hlavná 12/A, 040 01 Košice",                  # ADRESA + PSC + OBEC
    "Číslo OP: AB123456",                                        # CISLO_OP
    "Cestovný pas: SK1234567",                                   # CISLO_PASU
    "Vodičský preukaz č. KBXSWU0U",                              # VODICSKY_PREUKAZ
    "Telefón: 0905 123 456",                                     # TELEFON
    "Fax: 055 123 4567",                                         # FAX
    "E-mail: jan.novak@pravnik.sk",                              # EMAIL
    "",
    "Kupujúci: Ing. Mária Kováčová",                             # MENO (title + role anchor)
    "Rodné číslo: 905612/345",                                   # RODNE_CISLO 9-digit (the R0 bug)
    "Bydlisko: Štúrova 15, 811 06 Bratislava",                   # ADRESA + PSC + OBEC
    "Byt č. 12, vchod B, 3. poschodie, súp. č. 1234",            # CISLO_BYTU VCHOD POSCHODIE SUPISNE
    "Štátna príslušnosť: Slovenská republika",                   # STATNA_PRISLUSNOST
    "",
    "Článok II — Predmet zmluvy",
    "Nehnuteľnosť v k. ú. Michalovce, obec Michalovce,",         # KATASTER + OBEC
    "zapísaná na LV č. 1234, parc. č. 567/8.",                   # LV + PARCELA
    "Kúpna cena je 125 000,00 EUR.",                             # SUMA
    "",
    "Článok III — Platobné údaje",
    "IBAN: SK31 1200 0000 1987 4263 7541",                       # IBAN
    "Názov účtu: Ján Novák",                                     # NAZOV_UCTU
    "Kód banky 1200",                                            # KOD_BANKY
    "BIC: TATRSKBX",                                             # BIC
    "Číslo klienta: KL-99321",                                   # CISLO_KLIENTA
    "Účet: 013389-1543039117/1100",                              # BANKOVY_UCET
    "",
    "Článok IV — Ostatné dojednania",
    "Predávajúci je konateľom spoločnosti Stavby Východ s.r.o.,",  # ORG (not yet detected)
    "IČO: 36000001, DIČ: 2020123456, IČ DPH: SK2020123456,",     # ICO DIC IC_DPH
    "zapísanej v ORSR, Oddiel: Sro, Vložka č. 12345/V.",         # ORSR_VLOZKA
    "Vozidlo EČV KE-123AB, VIN 1HGBH41JXMN109186.",              # ECV + VIN
    "Spisová značka V-1234/2025.",                               # SPISOVA_ZNACKA
    "Svedok Peter Horváth potvrdzuje podpisy.",                  # MENO (role anchor)
    "Zmluvu pripravil Tomáš Zelinka pre klienta.",               # MENO (bare-name -> REVIEW)
    "",
    "Strana 1 z 1 — informácie na www.pravnik.sk",               # decoys: page footer + URL
]


def build_docx(path: Path) -> None:
    doc = Document()
    for line in PARAGRAPHS:
        doc.add_paragraph(line)
    # A name split across runs — the classic silent failure (context.md §10).
    p = doc.add_paragraph()
    for fragment in ("Podpísal: Ján ", "Nov", "ák, dňa ", "1.1.2025."):
        p.add_run(fragment)
    # PII in a table cell and in the header/footer.
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).paragraphs[0].add_run("Rodné číslo predávajúceho")
    table.cell(0, 1).paragraphs[0].add_run("850315/0018")
    table.cell(1, 0).paragraphs[0].add_run("Telefón kupujúceho")
    table.cell(1, 1).paragraphs[0].add_run("0911 402 917")
    section = doc.sections[0]
    section.header.paragraphs[0].add_run("Kúpna zmluva — Ján Novák")
    section.footer.paragraphs[0].add_run("Kontakt: maria.kovacova@email.sk")
    doc.core_properties.author = "JUDr. Ján Novák"
    doc.core_properties.company = "Stavby Východ s.r.o."
    doc.save(str(path))


def build_pdf(path: Path) -> None:
    """Draw the contract with an EMBEDDED Unicode font.

    The base-14 fonts `page.insert_text` defaults to have no Slovak glyphs, and PyMuPDF
    substitutes a placeholder rather than failing. The first version of this demo did exactly
    that, and the result was instructive: "Číslo klienta:" came out of the text layer as
    "·íslo klienta:", so the anchor-required CISLO_KLIENTA detector never fired and `KL-99321`
    survived into the redacted PDF. That was a defect in the FIXTURE, not in the detector — but
    a demo that quietly misrepresents the product is worse than no demo, so it uses the same
    embedded font the corpus builder uses for the same reason.
    """
    font_path = _find_font()
    font = fitz.Font(fontfile=font_path)
    doc = fitz.open()
    page = doc.new_page()
    writer = fitz.TextWriter(page.rect)
    y = 60
    for line in PARAGRAPHS:
        if line:
            writer.append((56, y), line, font=font, fontsize=9)
        y += 13
        if y > 780:
            writer.write_text(page)
            page = doc.new_page()
            writer = fitz.TextWriter(page.rect)
            y = 60
    writer.write_text(page)
    doc.set_metadata({"author": "JUDr. Ján Novák", "title": "Kúpna zmluva"})
    doc.subset_fonts()
    doc.save(str(path), garbage=4, deflate=True)
    doc.close()


def run_pipeline(src: Path) -> None:
    """Exactly what the GUI wizard does: scan → (reviewer ticks) → export."""
    print(f"\n{'=' * 78}\n  {src.name}\n{'=' * 78}")
    scan = scan_file(str(src), KNOWN_ENTITIES)
    if scan.error:
        print(f"  REFUSED: {scan.error}")
        return

    auto = [r for r in scan.rows if r.bucket == "auto"]
    review = [r for r in scan.rows if r.bucket == "review"]
    print(f"  review screen: {len(auto)} auto (pre-ticked), {len(review)} review (unticked)\n")
    print(f"  {'TYPE':<20} {'text':<34} {'checksum':<9} {'n':>3}  where")
    print(f"  {'-' * 20} {'-' * 34} {'-' * 9} {'-' * 3}  {'-' * 18}")
    for row in sorted(auto, key=lambda r: (r.type, r.text)):
        text = row.text if len(row.text) <= 33 else row.text[:30] + "..."
        print(f"  {row.type:<20} {text:<34} {row.checksum:<9} {row.count:>3}  "
              f"{', '.join(row.locations)[:18]}")
    if review:
        print("\n  REVIEW BUCKET (unticked — a human must decide):")
        for row in review:
            print(f"  {row.type:<20} {row.text:<34} {row.checksum:<9} {row.count:>3}")

    # The reviewer accepts every pre-ticked row and leaves the review bucket alone.
    checked = {r.group: (r.bucket == "auto") for r in scan.rows}
    decisions = build_decisions(scan.rows, checked, KNOWN_ENTITIES)
    out, report = export_file(str(src), KNOWN_ENTITIES, decisions)
    print(f"\n  exported: {Path(out).name}")
    print(f"  report:   {Path(report).name}")


def main() -> int:
    if DEMO_DIR.exists():
        shutil.rmtree(DEMO_DIR)
    DEMO_DIR.mkdir(parents=True)

    docx_path = DEMO_DIR / "kupna_zmluva_demo.docx"
    pdf_path = DEMO_DIR / "kupna_zmluva_demo.pdf"
    build_docx(docx_path)
    build_pdf(pdf_path)
    print(f"built {docx_path}  ({docx_path.stat().st_size:,} B)")
    print(f"built {pdf_path}  ({pdf_path.stat().st_size:,} B)")

    for src in (docx_path, pdf_path):
        run_pipeline(src)

    print(f"\n{'=' * 78}")
    print("  files in demo/:")
    for f in sorted(DEMO_DIR.iterdir()):
        print(f"    {f.name:<44} {f.stat().st_size:>9,} B")
    print(f"{'=' * 78}")
    print("  The two _report.txt files carry the SOURCE FORMAT in their names, so the .docx")
    print("  and .pdf reports of one document no longer overwrite each other (v1.1 Phase F).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
