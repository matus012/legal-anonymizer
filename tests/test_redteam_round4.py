"""Red-team round 4 — the code that landed tonight, and the two new refusals.

Every test here asserts THE PROPERTY THAT SHOULD HOLD. A test that asserted the bug would be
green now and turn red the day somebody fixed it, which is backwards. Findings that are not
fixed yet carry ``xfail(strict=True)``, so the day a fix lands the suite says so by XPASSing.

The tests that PASS are CONTROLS. Without them a refusal here would be indistinguishable from
a fixture that is simply malformed, and a "needle not found on any surface" would be
indistinguishable from a correct redaction (round 3 caught two of its own mistakes that way).

Findings are documented in ``redteam/FINDINGS_ROUND4.md``.
"""
from __future__ import annotations

import glob
import os
import time
import urllib.parse
import zipfile
from pathlib import Path

import fitz
import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from lxml import etree

from corpus.mutations import mutate_mojibake_cp1250
from detect.core import detect
from detect.normalize import normalize
from eval.extract import extract
from writer.docx_body import (_target_carries_pii, redact_docx_collect)
from writer.errors import UnreadableDocumentError
from writer.pdf_body import (NoTextLayerError, ShreddedTextLayerError,
                             UnreadableTextLayerError, page_has_unreadable_text,
                             page_is_shredded, redact_pdf)

REPO = Path(__file__).resolve().parents[1]
SYNTH = REPO / "data" / "synthetic"

# --------------------------------------------------------------------------------- PDF kit
# A hand-built PDF, because PyMuPDF's insert_text always writes a sane /ToUnicode map and
# three of this round's attacks are about what happens when the producer does not.


def _cmap(pairs):
    body = "\n".join(f"<{c:02X}> <{u:04X}>" for c, u in pairs)
    return (
        "/CIDInit /ProcSet findresource begin\n12 dict begin\nbegincmap\n"
        "/CIDSystemInfo <</Registry (Adobe) /Ordering (UCS) /Supplement 0>> def\n"
        "/CMapName /Adobe-Identity-UCS def\n/CMapType 2 def\n"
        "1 begincodespacerange\n<00> <FF>\nendcodespacerange\n"
        f"{len(pairs)} beginbfchar\n{body}\nendbfchar\nendcmap\nend\nend\n"
    ).encode("latin-1")


ASCII_MAP = [(c, c) for c in range(0x20, 0x7F)]


def build_pdf(lines, *, tounicode=None, fontsize=11.0):
    """``lines`` = [(x, y_from_top, raw_bytes)], one Tj each. The bytes go in verbatim, so the
    caller decides which encoding they are in."""
    objs, extra = {}, b"/Encoding /WinAnsiEncoding"
    if tounicode is not None:
        extra += b"/ToUnicode 6 0 R"
    content = []
    for x, y, raw in lines:
        esc = raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")
        content.append(b"BT /F1 " + f"{fontsize} Tf 1 0 0 1 {x} {842 - y} Tm".encode()
                       + b" (" + esc + b") Tj ET")
    stream = b"\n".join(content)
    objs[1] = b"<</Type/Catalog/Pages 2 0 R>>"
    objs[2] = b"<</Type/Pages/Kids[3 0 R]/Count 1>>"
    objs[3] = (b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 595 842]"
               b"/Resources<</Font<</F1 4 0 R>>>>/Contents 5 0 R>>")
    objs[4] = b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica " + extra + b">>"
    objs[5] = b"<</Length " + str(len(stream)).encode() + b">>stream\n" + stream + b"\nendstream"
    if tounicode is not None:
        cm = _cmap(tounicode)
        objs[6] = b"<</Length " + str(len(cm)).encode() + b">>stream\n" + cm + b"\nendstream"
    out, offsets = bytearray(b"%PDF-1.4\n"), {}
    for num in sorted(objs):
        offsets[num] = len(out)
        out += f"{num} 0 obj".encode() + objs[num] + b"endobj\n"
    xref_at, n = len(out), max(objs) + 1
    out += f"xref\n0 {n}\n".encode() + b"0000000000 65535 f \n"
    for num in range(1, n):
        out += (f"{offsets[num]:010d} 00000 n \n".encode() if num in offsets
                else b"0000000000 65535 f \n")
    out += f"trailer<</Size {n}/Root 1 0 R>>\nstartxref\n{xref_at}\n".encode() + b"%%EOF\n"
    return bytes(out)


SP = lambda s: " ".join(s)          # noqa: E731 — letter-spacing, razvetrené písmo
A = lambda s: s.encode("ascii")     # noqa: E731


def write_pdf(tmp_path, name, lines, **kw):
    p = tmp_path / f"{name}.pdf"
    p.write_bytes(build_pdf(lines, **kw))
    return str(p)


def pdf_text(path):
    d = fitz.open(path)
    try:
        return "\n".join(pg.get_text("text") for pg in d)
    finally:
        d.close()


# -------------------------------------------------------------------------------- DOCX kit
NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml" '
)


def doc_with_body(fragments, path):
    d = Document()
    body = d.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    sect = body.find(qn("w:sectPr"))
    for frag in fragments:
        for child in list(parse_xml(f"<w:_wrap {NS}>{frag}</w:_wrap>")):
            sect.addprevious(child) if sect is not None else body.append(child)
    d.save(str(path))
    return str(path)


def patch_package(src, dst, *, add=None, replace=None, content_types=None,
                  extra_rels=None, doc_rels=None):
    add, replace = add or {}, replace or {}
    with zipfile.ZipFile(src) as z:
        items = [(n, z.read(n)) for n in z.namelist()]
    out = []
    for name, data in items:
        data = replace.get(name, data)
        if name == "[Content_Types].xml" and content_types:
            data = data.replace(b"</Types>", ("".join(content_types) + "</Types>").encode())
        if name == "_rels/.rels" and extra_rels:
            data = data.replace(b"</Relationships>",
                                ("".join(extra_rels) + "</Relationships>").encode())
        if name == "word/_rels/document.xml.rels" and doc_rels:
            data = data.replace(b"</Relationships>",
                                ("".join(doc_rels) + "</Relationships>").encode())
        out.append((name, data))
    have = {n for n, _ in out}
    out += [(n, d) for n, d in add.items() if n not in have]
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in out:
            z.writestr(name, data)
    return str(dst)


def textbox(text):
    return ('<w:r><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r>'
            f'<w:t xml:space="preserve">{text}</w:t></w:r></w:p></w:txbxContent>'
            '</v:textbox></v:shape></w:pict></w:r>')


def doc_text(path):
    x = zipfile.ZipFile(path).read("word/document.xml")
    return "".join(etree.fromstring(x).itertext())


def doc_xml(path):
    return zipfile.ZipFile(path).read("word/document.xml").decode()


def reopens(path):
    """Word's own schema terms, as far as this repo can check them without Word: the package
    re-opens through python-docx (so every part it needs parses and every relationship it
    follows resolves) and the part list is intact. A fixture Word would reject proves nothing,
    and neither does an output it would reject."""
    assert Document(path).element.body is not None
    names = set(zipfile.ZipFile(path).namelist())
    assert {"[Content_Types].xml", "_rels/.rels", "word/document.xml",
            "word/_rels/document.xml.rels"} <= names
    for name in names:
        if name.endswith(".xml") or name.endswith(".rels"):
            etree.fromstring(zipfile.ZipFile(path).read(name))
    return True


def redact(tmp_path, name, fragments, known=(), patch=None):
    src = doc_with_body(fragments, tmp_path / f"{name}.docx")
    if patch is not None:
        src = patch(src, str(tmp_path / f"{name}_p.docx"))
    out = str(tmp_path / f"{name}_r.docx")
    lm = redact_docx_collect(src, out, list(known))
    return out, lm


def surfaces_with(path, needle):
    return sorted(s for s, t in extract(path).by_surface.items() if needle in t)


KNOWN = ["Ján Novák", "Mária Kováčová"]


# ==================================================================== P0 / controls: refusals
def test_control_no_corpus_pdf_is_refused():
    """P0. Amendment 16 claims the thresholds were tuned for ZERO false refusals, 0 of 71.
    Re-measured here: if this fails, every number in group P is unattributable."""
    hits = []
    for f in sorted(glob.glob(str(SYNTH / "*.pdf"))):
        d = fitz.open(f)
        for i, page in enumerate(d, 1):
            t = page.get_text("text")
            if page_is_shredded(t) or page_has_unreadable_text(t):
                hits.append((os.path.basename(f), i))
        d.close()
    assert hits == [], f"corpus PDFs trip the new refusals: {hits}"


def test_control_unspaced_title_page_is_accepted(tmp_path):
    """P1's control: the SAME notarial title page without letter-spacing must be processed."""
    src = write_pdf(tmp_path, "ctl", [
        (120, 120, A("NOTARSKA ZAPISNICA")),
        (120, 160, A("spisana dna 17. septembra 2026 pred notarom v Kosiciach")),
        (120, 180, A("N 245/2026, Nz 33214/2026, kancelaria Hlavna 25"))])
    redact_pdf(src, str(tmp_path / "ctl_r.pdf"), [])


def test_control_two_symbol_bullets_are_accepted(tmp_path):
    """P3's control: TWO Symbol bullets are under ``_UNDECODABLE_MAX`` and must be accepted —
    which is what makes the three-bullet refusal a threshold effect and not a broken fixture."""
    src = write_pdf(tmp_path, "b2", [
        (80, 100, A("Predavajuci sa zavazuje:")),
        (90, 130, b"\xb7 odovzdat nehnutelnost kupujucemu"),
        (90, 150, b"\xb7 predlozit list vlastnictva"),
        (90, 170, A("- uhradit spravny poplatok")),
        (80, 200, A("Kupujuci: Jan Novak, rodne cislo 850101/1234"))],
        tounicode=ASCII_MAP + [(0xB7, 0xF0B7)])
    out = str(tmp_path / "b2_r.pdf")
    redact_pdf(src, out, [])
    assert "Jan Novak" not in pdf_text(out)


# ============================================================ P1 / P3: FALSE REFUSALS
@pytest.mark.xfail(reason="R4-P1: a letter-spaced Slovak notarial title page is refused",
                   strict=True)
def test_letterspaced_notarial_title_page_is_not_refused(tmp_path):
    src = write_pdf(tmp_path, "p1", [
        (120, 120, A(SP("NOTARSKA") + "  " + SP("ZAPISNICA"))),
        (120, 160, A("spisana dna 17. septembra 2026")),
        (120, 180, A("N 245/2026, Nz 33214/2026"))])
    redact_pdf(src, str(tmp_path / "p1_r.pdf"), [])


@pytest.mark.xfail(reason="R4-P1b: the standard Slovak judgment heading is refused",
                   strict=True)
def test_letterspaced_judgment_heading_is_not_refused(tmp_path):
    src = write_pdf(tmp_path, "p1b", [
        (120, 120, A(SP("ROZSUDOK"))),
        (100, 150, A(SP("V") + " " + SP("MENE") + " " + SP("SLOVENSKEJ") + " "
                     + SP("REPUBLIKY"))),
        (120, 200, A("Okresny sud Kosice I, sp. zn. 12Cb/345/2025"))])
    redact_pdf(src, str(tmp_path / "p1b_r.pdf"), [])


# FIXED 2026-09-17, overnight run. The marker is gone rather than flipped to xpass: a finding that has been
# fixed must become an ordinary regression test, or a later regression puts it back to
# "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
def test_symbol_bullet_list_is_not_refused(tmp_path):
    src = write_pdf(tmp_path, "p3", [
        (80, 100, A("Predavajuci sa zavazuje:")),
        (90, 130, b"\xb7 odovzdat nehnutelnost kupujucemu"),
        (90, 150, b"\xb7 predlozit list vlastnictva"),
        (90, 170, b"\xb7 uhradit spravny poplatok"),
        (80, 200, A("Kupujuci: Jan Novak, rodne cislo 850101/1234"))],
        tounicode=ASCII_MAP + [(0xB7, 0xF0B7)])
    redact_pdf(src, str(tmp_path / "p3_r.pdf"), [])


# FIXED 2026-09-17, overnight run. The marker is gone rather than flipped to xpass: a finding that has been
# fixed must become an ordinary regression test, or a later regression puts it back to
# "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
def test_symbol_bullets_do_not_refuse_the_corpus():
    """The same three glyphs added to real corpus pages. 15 documents, not 71, to keep the
    dev loop quick; the full sweep is in the findings file."""
    refused = 0
    files = sorted(glob.glob(str(SYNTH / "*.pdf")))[:15]
    for f in files:
        d = fitz.open(f)
        if any(page_has_unreadable_text("\uf0b7\uf0b7\uf0b7" + pg.get_text("text")) for pg in d):
            refused += 1
        d.close()
    assert refused == 0, f"{refused}/{len(files)} ordinary documents refused over three bullets"


# ================================================ I1 / I2 / I3: are the round-3 fixes complete?
PROSE = [
    (72, 90, A("KUPNA ZMLUVA uzavreta podla par. 588 Obcianskeho zakonnika")),
    (72, 120, A("Zmluvne strany sa dohodli na prevode nehnutelnosti zapisanej na")),
    (72, 140, A("liste vlastnictva vedenom Okresnym uradom Kosice, katastralny odbor.")),
    (72, 220, A("Kupujuci sa zavazuje uhradit kupnu cenu v lehote tridsiatich dni odo")),
    (72, 240, A("dna podpisu tejto zmluvy na ucet uvedeny v clanku IV tejto zmluvy.")),
    (72, 270, A("Zmluva nadobuda ucinnost dnom vkladu vlastnickeho prava do katastra.")),
]


def _i1_pages(spaced: bool):
    if spaced:
        party = [(72, 170, A("Predavajuci:  " + SP("Jan") + "   " + SP("Novak"))),
                 (72, 190, A("rodne cislo " + SP("850101/1234")))]
    else:
        party = [(72, 170, A("Predavajuci:  Jan Novak")),
                 (72, 190, A("rodne cislo 850101/1234"))]
    return PROSE[:3] + party + PROSE[3:]


def test_control_unspaced_party_line_is_redacted(tmp_path):
    """I1's control. Identical page, party line NOT letter-spaced: both surfaces must go."""
    src = write_pdf(tmp_path, "i1c", _i1_pages(spaced=False))
    out = str(tmp_path / "i1c_r.pdf")
    redact_pdf(src, out, ["Jan Novak"])
    t = pdf_text(out)
    assert "Jan Novak" not in t and "850101/1234" not in t
    assert "[MENO_1]" in t and "[RODNE_CISLO_1]" in t


@pytest.mark.xfail(reason="R4-I1: one letter-spaced party line stays under the page-level "
                          "shred ratio, so the name and the rodne cislo survive and the leak "
                          "gate scores the output clean", strict=True)
def test_partially_letterspaced_party_line_is_redacted(tmp_path):
    src = write_pdf(tmp_path, "i1", _i1_pages(spaced=True))
    out = str(tmp_path / "i1_r.pdf")
    try:
        redact_pdf(src, out, ["Jan Novak"])
    except NoTextLayerError:
        return  # a refusal would also be an acceptable outcome; today it is neither
    t = pdf_text(out)
    assert "J a n   N o v a k" not in t, "the party's name is still drawn on the page"
    assert "8 5 0 1 0 1 / 1 2 3 4" not in t, "the rodne cislo is still drawn on the page"


def test_control_subset_font_free_page_is_redacted(tmp_path):
    """I2's control: the same page with no subset-font glyphs redacts the name."""
    src = write_pdf(tmp_path, "i2c", [
        (72, 90, A("Predavajuci: Maria Kovacova, bytom Hlavna 25, Kosice")),
        (72, 120, A("Kupna cena bola uhradena v plnej vyske pri podpise tejto zmluvy.")),
        (72, 150, A("Zmluva nadobuda ucinnost dnom vkladu do katastra nehnutelnosti."))],
        tounicode=ASCII_MAP)
    out = str(tmp_path / "i2c_r.pdf")
    redact_pdf(src, out, ["Maria Kovacova"])
    assert "Maria Kovacova" not in pdf_text(out)


# FIXED 2026-09-17, overnight run. The marker is gone rather than flipped to xpass: a finding that has been
# fixed must become an ordinary regression test, or a later regression puts it back to
# "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
def test_two_undecodable_glyphs_do_not_hide_a_name(tmp_path):
    src = write_pdf(tmp_path, "i2", [
        (72, 90, b"Predavajuci: M\x01ria Kov\x02cova, bytom Hlavna 25, Kosice"),
        (72, 120, A("Kupna cena bola uhradena v plnej vyske pri podpise tejto zmluvy.")),
        (72, 150, A("Zmluva nadobuda ucinnost dnom vkladu do katastra nehnutelnosti."))],
        tounicode=ASCII_MAP + [(0x01, 0x0001), (0x02, 0x0002)])
    out = str(tmp_path / "i2_r.pdf")
    try:
        redact_pdf(src, out, ["Maria Kovacova"])
    except NoTextLayerError:
        return
    assert "M\x01ria Kov\x02cova" not in pdf_text(out)


def test_control_mojibake_mutation_is_character_local():
    """The round-4 mutation must be 1:1 per character — the homomorphism requirement
    corpus/mutations.py states. If it is not, every mojibake number is a harness artefact."""
    for s in ["Ján Kováč", "Ľubomír Ďurčo", "SK31 1200 0000", "850101/1234"]:
        assert len(mutate_mojibake_cp1250(s)) == len(s)
    assert mutate_mojibake_cp1250("Ján Kováč") == "Ján Kováè"
    assert mutate_mojibake_cp1250("850101/1234") == "850101/1234"


# FIXED 2026-09-17, daytime run. The marker is gone rather than flipped to xpass: a finding that has
# been fixed must become an ordinary regression test, or a later regression puts it back to
# "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
# detect/normalize.py folds the cp1250-as-WinAnsi rewrite back, 1:1 and offset-preservingly,
# in any unit that shows evidence of the failure; the party is found and the ORIGINAL mojibake
# surface is what the writer redacts off the page.
def test_mojibake_text_layer_still_finds_the_party(tmp_path):
    src = write_pdf(tmp_path, "i3", [
        (72, 90, "Predavajuci: Ľubomír Ďurčo, bytom Hlavna 25, Kosice".encode("cp1250")),
        (72, 120, A("Kupna cena bola uhradena v plnej vyske pri podpise tejto zmluvy.")),
        (72, 150, A("Zmluva nadobuda ucinnost dnom vkladu do katastra nehnutelnosti."))])
    out = str(tmp_path / "i3_r.pdf")
    try:
        redact_pdf(src, out, ["Ľubomír Ďurčo"])
    except NoTextLayerError:
        return
    assert "¼ubomír Ïurèo" not in pdf_text(out)


# ================================================================ T: the rewritten traversal
def test_control_smarttag_run_is_redacted(tmp_path):
    """T6 — round 3 suspected the `.//w:r` fix would cover `<w:smartTag>` and could not
    confirm it. Confirmed here."""
    out, _ = redact(tmp_path, "t6", [
        '<w:p><w:r><w:t xml:space="preserve">Predávajúci: </w:t></w:r>'
        '<w:smartTag w:uri="urn:schemas-microsoft-com:office:smarttags" w:element="PersonName">'
        '<w:r><w:t xml:space="preserve">Ján Novák</w:t></w:r></w:smartTag>'
        '<w:r><w:t xml:space="preserve"> prehlasuje.</w:t></w:r></w:p>'], KNOWN)
    assert surfaces_with(out, "Ján Novák") == []


def test_control_sdt_around_a_table_is_redacted(tmp_path):
    """T3 — a content control wrapping a whole TABLE. Clean, as the rewrite intends."""
    out, _ = redact(tmp_path, "t3", [
        '<w:sdt><w:sdtPr><w:alias w:val="Strany"/></w:sdtPr><w:sdtContent><w:tbl><w:tr><w:tc>'
        '<w:p><w:r><w:t xml:space="preserve">Predávajúci: Ján Novák</w:t></w:r></w:p>'
        '</w:tc></w:tr></w:tbl></w:sdtContent></w:sdt>'], KNOWN)
    assert surfaces_with(out, "Ján Novák") == []


def test_control_fully_covered_hyperlink_run_keeps_its_label(tmp_path):
    """T7 — a hyperlink whose whole run is the PII. The address goes and `[EMAIL_1]` stays
    inside the hyperlink, so the reviewer can still see what was removed."""
    out, _ = redact(tmp_path, "t7", [
        '<w:p><w:r><w:t xml:space="preserve">Kontakt: </w:t></w:r>'
        '<w:hyperlink r:id="rIdX"><w:r><w:t xml:space="preserve">jan.novak@advokat.sk</w:t>'
        '</w:r></w:hyperlink><w:r><w:t xml:space="preserve"> (kancelária)</w:t></w:r></w:p>'])
    assert surfaces_with(out, "jan.novak@advokat.sk") == []
    assert "[EMAIL_1]" in doc_text(out)


def test_control_textbox_in_an_empty_anchor_paragraph_is_not_welded(tmp_path):
    """T9's control, and the shape the corpus actually contains: when the anchor paragraph
    carries no text of its own there is nothing to weld, and the textbox redacts cleanly."""
    out, lm = redact(tmp_path, "t9c", [
        f'<w:p>{textbox("Predávajúci: Ján Novák")}</w:p>'], KNOWN)
    assert surfaces_with(out, "Ján Novák") == []
    assert doc_text(out) == "Predávajúci: [MENO_1]"


def test_control_linked_header_is_one_shared_element():
    """T2's premise, measured. `roots` appends `section.header._element` once per section and
    a linked header resolves to the SAME element, so those paragraphs ARE walked twice. The
    report is not double-counted (the second pass detects nothing in `[MENO_1]`) — recorded
    because the premise holds and only the consequence did not."""
    d = Document()
    d.add_paragraph("A")
    d.add_section(WD_SECTION.NEW_PAGE)
    assert d.sections[1].header.is_linked_to_previous
    assert d.sections[0].header._element is d.sections[1].header._element


# FIXED 2026-09-17, overnight run. The marker is gone rather than flipped to xpass: a finding that has been
# fixed must become an ordinary regression test, or a later regression puts it back to
# "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
def test_anchored_textbox_text_is_not_welded_onto_its_paragraph(tmp_path):
    out, lm = redact(tmp_path, "t9b", [
        '<w:p><w:r><w:t xml:space="preserve">Predávajúci: Ján Novák</w:t></w:r>'
        + textbox("Kupujúci: Mária Kováčová") + '</w:p>'], KNOWN)
    assert "Kupujúci" in doc_text(out), "the other party's role label was deleted"
    recorded = [s for occ in lm.occurrences.values() for _, s in occ]
    assert "Ján NovákKupujúci" not in recorded, \
        "the report records a surface that exists in no part of the source document"


# FIXED 2026-09-17, overnight run. The marker is gone rather than flipped to xpass: a finding that has been
# fixed must become an ordinary regression test, or a later regression puts it back to
# "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
def test_welded_span_does_not_leave_half_a_name_behind(tmp_path):
    out, _ = redact(tmp_path, "t9d", [
        '<w:p><w:r><w:t xml:space="preserve">Podpísal Ján</w:t></w:r>'
        + textbox("Novák, predávajúci") + '</w:p>'], KNOWN)
    assert "Ján" not in doc_text(out)


def _alternate_content(text):
    tb = ('<w:txbxContent><w:p><w:r><w:t xml:space="preserve">' + text
          + "</w:t></w:r></w:p></w:txbxContent>")
    return ('<w:p><w:r><mc:AlternateContent>'
            '<mc:Choice Requires="wps"><w:drawing><wp:inline><a:graphic><a:graphicData>'
            f'<wps:wsp><wps:txbx>{tb}</wps:txbx></wps:wsp>'
            '</a:graphicData></a:graphic></wp:inline></w:drawing></mc:Choice>'
            f'<mc:Fallback><w:pict><v:shape><v:textbox>{tb}</v:textbox></v:shape></w:pict>'
            '</mc:Fallback></mc:AlternateContent></w:r></w:p>')


def _mc_branches(path):
    """The <mc:Choice> and <mc:Fallback> halves of word/document.xml, as text."""
    x = zipfile.ZipFile(path).read("word/document.xml").decode()
    return (x[x.index("<mc:Choice"):x.index("</mc:Choice>")],
            x[x.index("<mc:Fallback>"):x.index("</mc:Fallback>")])


def test_control_alternate_content_needle_is_reachable_in_both_branches(tmp_path):
    """T1's fixture control, and the one the finding most needed. python-docx cannot emit an
    <mc:AlternateContent> at all, so the fixture is raw OOXML and a "needle not found" would
    otherwise be indistinguishable from a fixture Word would reject.

    Asserted on the UNREDACTED package: the name really is in BOTH branches before the writer
    runs, so a test that finds it gone from both is measuring the writer and not the fixture.
    The package is re-opened through python-docx as well, so a fixture that does not parse
    fails here rather than silently passing the redaction tests."""
    src = doc_with_body([_alternate_content("Predávajúci: Ján Novák")], tmp_path / "t1c.docx")
    choice, fallback = _mc_branches(src)
    assert "Ján Novák" in choice and "Ján Novák" in fallback
    assert Document(src).element.body is not None
    assert "word/document.xml" in zipfile.ZipFile(src).namelist()


def test_control_a_one_branch_scrub_is_visibly_a_leak(tmp_path):
    """The other half of T1's control: what a writer that scrubbed only <mc:Choice> would
    leave. Hand-built as the output such a writer would produce -- Choice already reading
    `[MENO_1]`, Fallback still reading the name -- and put through the leak extractor.

    Without this, `test_alternate_content_fallback_copy_is_not_damaged` could pass because the
    Fallback branch is invisible to `extract()` rather than because the writer scrubbed it."""
    half = _alternate_content("Predávajúci: Ján Novák").replace(
        "Predávajúci: Ján Novák", "Predávajúci: [MENO_1]", 1)      # only the Choice copy
    src = doc_with_body([half], tmp_path / "t1half.docx")
    choice, fallback = _mc_branches(src)
    assert "[MENO_1]" in choice and "Ján Novák" in fallback
    assert surfaces_with(src, "Ján Novák") != [], \
        "a name left in the Fallback branch must be gradeable, or T1 proves nothing"


# FIXED 2026-09-17, daytime run. The marker is gone rather than flipped to xpass: a finding
# that has been fixed must become an ordinary regression test, or a later regression puts it
# back to "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
def test_alternate_content_fallback_copy_is_not_damaged(tmp_path):
    out, lm = redact(tmp_path, "t1", [_alternate_content("Predávajúci: Ján Novák")], KNOWN)
    choice, fallback = _mc_branches(out)
    assert "Predávajúci" in fallback, "the label was deleted from the Fallback copy"
    # BOTH copies scrubbed, never one: a document scrubbed in one branch only renders clean in
    # whichever Word branch you happen to open it in, and leaks in the other.
    assert "Ján Novák" not in choice and "Ján Novák" not in fallback
    assert "[MENO_1]" in choice and "[MENO_1]" in fallback
    assert surfaces_with(out, "Ján Novák") == []
    assert sum(len(v) for v in lm.occurrences.values()) == 1, \
        "one name on the page must be one occurrence in the report"


T4_FRAGMENT = ('<w:p><w:fldSimple w:instr=" HYPERLINK &quot;mailto:jan.novak@advokat.sk&quot; ">'
               '<w:r><w:t xml:space="preserve">kontakt na advokáta</w:t></w:r>'
               "</w:fldSimple></w:p>")


def test_control_fldsimple_instruction_is_reachable(tmp_path):
    """T4/T5's fixture control. python-docx emits neither field spelling, so the fixtures are
    raw OOXML; this asserts the package parses, keeps its part list, and that the address
    really is on a gradeable surface BEFORE the writer runs -- otherwise "not found" after
    redaction would be indistinguishable from a malformed fixture."""
    src = doc_with_body([T4_FRAGMENT], tmp_path / "t4c.docx")
    assert Document(src).element.body is not None
    assert "word/_rels/document.xml.rels" in zipfile.ZipFile(src).namelist()
    assert surfaces_with(src, "jan.novak@advokat.sk") != []


# FIXED 2026-09-17, daytime run. The marker is gone rather than flipped to xpass: a finding
# that has been fixed must become an ordinary regression test, or a later regression puts it
# back to "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
def test_fldsimple_hyperlink_destination_is_scrubbed(tmp_path):
    out, _ = redact(tmp_path, "t4", [T4_FRAGMENT])
    assert surfaces_with(out, "jan.novak@advokat.sk") == []


# FIXED 2026-09-17, daytime run. The marker is gone rather than flipped to xpass: a finding
# that has been fixed must become an ordinary regression test, or a later regression puts it
# back to "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
def test_instrtext_hyperlink_destination_is_scrubbed(tmp_path):
    out, _ = redact(tmp_path, "t5", [
        '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText xml:space="preserve"> HYPERLINK '
        '"mailto:jan.novak@advokat.sk" </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        '<w:r><w:t xml:space="preserve">Ján Novák</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'], KNOWN)
    assert surfaces_with(out, "jan.novak@advokat.sk") == []


def test_instrtext_split_across_runs_is_still_scrubbed(tmp_path):
    """Word splits one field instruction across several <w:instrText> runs whenever it feels
    like it, so a per-element scrub would see half a URL and find nothing in it. The fixture
    cuts the address in two, mid-local-part, which is the shape an RTF->DOCX conversion
    actually produces."""
    out, _ = redact(tmp_path, "t5b", [
        '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText xml:space="preserve"> HYPERLINK "mailto:jan.no</w:instrText></w:r>'
        '<w:r><w:instrText xml:space="preserve">vak@advokat.sk" </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        '<w:r><w:t xml:space="preserve">Ján Novák</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'], KNOWN)
    assert surfaces_with(out, "jan.novak@advokat.sk") == []


def test_a_field_code_with_an_innocuous_destination_survives(tmp_path):
    """The quiet half, and the one this project has broken twice: a field-code link to the
    statute book is not PII and must still work after redaction."""
    out, _ = redact(tmp_path, "t4ok", [
        '<w:p><w:fldSimple w:instr=" HYPERLINK &quot;https://www.slov-lex.sk/&quot; ">'
        '<w:r><w:t xml:space="preserve">Zbierka zákonov</w:t></w:r></w:fldSimple></w:p>'])
    assert "https://www.slov-lex.sk/" in doc_xml(out)
    assert "removed.invalid" not in doc_xml(out)


# FIXED 2026-09-17, daytime run. The marker is gone rather than flipped to xpass: a finding
# that has been fixed must become an ordinary regression test, or a later regression puts it
# back to "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
def test_databound_content_control_store_is_scrubbed(tmp_path):
    store = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             '<zmluva xmlns="http://firma.sk/zmluva"><predavajuci>Ján Novák</predavajuci>'
             "<rodneCislo>850101/1234</rodneCislo></zmluva>").encode()
    out, _ = redact(tmp_path, "t8", [
        '<w:p><w:r><w:t xml:space="preserve">Predávajúci: </w:t></w:r></w:p>'
        '<w:sdt><w:sdtPr><w:alias w:val="predavajuci"/>'
        '<w:dataBinding w:xpath="/ns0:zmluva[1]/ns0:predavajuci[1]" '
        'w:storeItemID="{11111111-2222-3333-4444-555555555555}"/></w:sdtPr><w:sdtContent>'
        '<w:p><w:r><w:t xml:space="preserve">Ján Novák</w:t></w:r></w:p></w:sdtContent></w:sdt>'],
        KNOWN, patch=lambda a, b: patch_package(a, b,
                                                replace={"customXml/item1.xml": store}))
    assert surfaces_with(out, "Ján Novák") == []
    assert surfaces_with(out, "850101/1234") == []
    # The binding is gone as well as the store, so nothing re-populates the control if the
    # store is ever restored from a backup of the original.
    assert "dataBinding" not in doc_xml(out)
    assert reopens(out)


def test_control_databound_store_is_a_gradeable_surface(tmp_path):
    """T8's fixture control. The custom XML part is what Word actually displays, and it is a
    part python-docx copies through byte-for-byte -- so if `extract()` could not read it, the
    test above would pass on a fixture that never carried the needle at all."""
    store = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             '<zmluva xmlns="http://firma.sk/zmluva"><predavajuci>Ján Novák</predavajuci>'
             "<rodneCislo>850101/1234</rodneCislo></zmluva>").encode()
    src = doc_with_body(['<w:p><w:r><w:t>x</w:t></w:r></w:p>'], tmp_path / "t8c.docx")
    src = patch_package(src, str(tmp_path / "t8c_p.docx"),
                        replace={"customXml/item1.xml": store})
    assert reopens(src)
    assert surfaces_with(src, "Ján Novák") != []
    assert surfaces_with(src, "850101/1234") != []


# ========================================================= M: package parts nobody attacked
# FIXED 2026-09-17, overnight run. The marker is gone rather than flipped to xpass: a finding that has been
# fixed must become an ordinary regression test, or a later regression puts it back to
# "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
def test_settings_docvars_are_scrubbed(tmp_path):
    def patch(a, b):
        with zipfile.ZipFile(a) as z:
            s = z.read("word/settings.xml").decode()
        s = s.replace("</w:settings>",
                      '<w:docVars><w:docVar w:name="ClientName" w:val="Ján Novák"/>'
                      '<w:docVar w:name="ClientRC" w:val="850101/1234"/></w:docVars>'
                      "</w:settings>")
        return patch_package(a, b, replace={"word/settings.xml": s.encode()})

    out, _ = redact(tmp_path, "m1", [
        '<w:p><w:r><w:t xml:space="preserve">Predávajúci: Ján Novák prehlasuje.</w:t></w:r>'
        '</w:p>'], KNOWN, patch=patch)
    assert surfaces_with(out, "Ján Novák") == []
    assert surfaces_with(out, "850101/1234") == []


# FIXED 2026-09-17, overnight run. The marker is gone rather than flipped to xpass: a finding that has been
# fixed must become an ordinary regression test, or a later regression puts it back to
# "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
def test_docprops_custom_properties_are_scrubbed(tmp_path):
    props = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Properties '
             'xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
             'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
             '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="Klient">'
             "<vt:lpwstr>Ján Novák</vt:lpwstr></property></Properties>").encode()

    def patch(a, b):
        return patch_package(
            a, b, add={"docProps/custom.xml": props},
            content_types=['<Override PartName="/docProps/custom.xml" ContentType="'
                           'application/vnd.openxmlformats-officedocument.custom-properties'
                           '+xml"/>'],
            extra_rels=['<Relationship Id="rIdCust" Type="http://schemas.openxmlformats.org/'
                        'officeDocument/2006/relationships/custom-properties" '
                        'Target="docProps/custom.xml"/>'])

    out, _ = redact(tmp_path, "m2", [
        '<w:p><w:r><w:t xml:space="preserve">Predávajúci: Ján Novák prehlasuje.</w:t></w:r>'
        '</w:p>'], KNOWN, patch=patch)
    assert surfaces_with(out, "Ján Novák") == []


# FIXED 2026-09-17, overnight run. The marker is gone rather than flipped to xpass: a finding that has been
# fixed must become an ordinary regression test, or a later regression puts it back to
# "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
def test_people_xml_authors_are_scrubbed(tmp_path):
    people = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w15:people '
              'xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml" '
              'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
              '<w15:person w15:author="JUDr. Ján Novák"><w15:presenceInfo w15:providerId="AD"'
              ' w15:userId="jan.novak@advokat.sk"/></w15:person></w15:people>').encode()

    def patch(a, b):
        return patch_package(
            a, b, add={"word/people.xml": people},
            content_types=['<Override PartName="/word/people.xml" ContentType="application/'
                           'vnd.openxmlformats-officedocument.wordprocessingml.people+xml"/>'],
            doc_rels=['<Relationship Id="rIdPeople" Type="http://schemas.microsoft.com/office/'
                      '2011/relationships/people" Target="people.xml"/>'])

    out, _ = redact(tmp_path, "m3", [
        '<w:p><w:r><w:t xml:space="preserve">Predávajúci: Ján Novák prehlasuje.</w:t></w:r>'
        '</w:p>'], KNOWN, patch=patch)
    assert surfaces_with(out, "Ján Novák") == []
    assert surfaces_with(out, "jan.novak@advokat.sk") == []


MOVE_DOC = (
    '<w:p><w:moveFrom w:id="3" w:author="JUDr. Ján Novák" w:date="2026-01-01T00:00:00Z">'
    '<w:r><w:delText xml:space="preserve">Kupujúci: Mária Kováčová, rodné číslo 855612/7788'
    "</w:delText></w:r></w:moveFrom></w:p>"
    '<w:p><w:moveTo w:id="4" w:author="JUDr. Ján Novák" w:date="2026-01-01T00:00:00Z">'
    '<w:r><w:t xml:space="preserve">Kupujúci: Mária Kováčová</w:t></w:r></w:moveTo></w:p>')


# FIXED 2026-09-17, overnight run. The marker is gone rather than flipped to xpass: a finding that has been
# fixed must become an ordinary regression test, or a later regression puts it back to
# "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
def test_movefrom_text_is_removed(tmp_path):
    out, _ = redact(tmp_path, "m4a", [MOVE_DOC], KNOWN)
    assert surfaces_with(out, "855612/7788") == []


# FIXED 2026-09-17, daytime run. The marker is gone rather than flipped to xpass: a finding
# that has been fixed must become an ordinary regression test, or a later regression puts it
# back to "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
def test_revision_authors_outside_ins_and_del_are_scrubbed(tmp_path):
    out, _ = redact(tmp_path, "m4b", [
        '<w:p><w:pPr><w:pPrChange w:id="1" w:author="JUDr. Ján Novák" '
        'w:date="2026-01-01T00:00:00Z"><w:pPr/></w:pPrChange></w:pPr>'
        '<w:r><w:rPr><w:rPrChange w:id="2" w:author="Mária Kováčová" '
        'w:date="2026-01-01T00:00:00Z"><w:rPr/></w:rPrChange></w:rPr>'
        '<w:t xml:space="preserve">Text odseku.</w:t></w:r></w:p>'], KNOWN)
    assert surfaces_with(out, "Ján Novák") == []
    assert surfaces_with(out, "Mária Kováčová") == []
    # The revision elements themselves are NOT removed -- they carry only the previous
    # formatting -- so the document still round-trips.
    assert reopens(out)
    assert "Text odseku." in doc_text(out)


def test_control_revision_author_is_a_gradeable_surface(tmp_path):
    """M4b's fixture control: a w:author on a w:pPrChange really is read by `extract()` (as
    xml_attributes), so "not found" above is a scrub and not a blind spot."""
    src = doc_with_body([
        '<w:p><w:pPr><w:pPrChange w:id="1" w:author="JUDr. Ján Novák" '
        'w:date="2026-01-01T00:00:00Z"><w:pPr/></w:pPrChange></w:pPr>'
        '<w:r><w:t xml:space="preserve">Text odseku.</w:t></w:r></w:p>'],
        tmp_path / "m4bc.docx")
    assert reopens(src)
    assert surfaces_with(src, "Ján Novák") == ["xml_attributes"]


# ============================================================== R: _scrub_rel_targets
REAL_LINKS = [
    "https://www.slov-lex.sk/pravne-predpisy/SK/ZZ/2016/18/20180101",
    "https://www.justice.gov.sk/Stranky/Sudy/Zoznam-sudov.aspx",
    "https://www.mfsr.sk/sk/dane-cla-ucto/",
]


# FIXED 2026-09-17, overnight run. The marker is gone rather than flipped to xpass: a finding that has been
# fixed must become an ordinary regression test, or a later regression puts it back to
# "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
@pytest.mark.parametrize("url", REAL_LINKS)
def test_real_slovak_legal_links_survive_the_scrub(url):
    assert not _target_carries_pii(urllib.parse.unquote(url), [], None)


def test_control_a_pii_bearing_target_is_still_scrubbed():
    """R1's control: the rule must still fire on the target it exists for."""
    assert _target_carries_pii(
        urllib.parse.unquote("https://dms.firma.sk/klienti/Jan%20Novak/zmluva.pdf"), [], None)
    assert _target_carries_pii("mailto:jan.novak@advokat.sk", [], None)


# FIXED 2026-09-17, overnight run. The marker is gone rather than flipped to xpass: a finding that has been
# fixed must become an ordinary regression test, or a later regression puts it back to
# "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
def test_double_encoded_target_is_scrubbed():
    t = "https://dms.firma.sk/klienti/Jan%2520Novak/zmluva.pdf"
    assert _target_carries_pii(urllib.parse.unquote(t), [], None)


# FIXED 2026-09-17, daytime run. The marker is gone rather than flipped to xpass: a finding
# that has been fixed must become an ordinary regression test, or a later regression puts it
# back to "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
@pytest.mark.parametrize("target", [
    "file:///C:/Users/jan.novak/Documents/zmluva.docx",
    r"\\fileserver\users\jan.novak\matter\zmluva.docx",
    "file:///C:/Users/Jan.Novak/Documents/zmluva.docx",      # the round's own control
])
def test_lowercase_username_in_a_file_target_is_scrubbed(target):
    assert _target_carries_pii(urllib.parse.unquote(target), [], None)


def test_a_filesystem_target_with_no_account_name_survives():
    """The narrowing that keeps the R3 fix from becoming the R1 defect: the rule fires on a
    dotted DIRECTORY segment, not on the file name. A share path with no account name in it is
    not personal data and is left alone."""
    assert not _target_carries_pii("file:///C:/Vzory/kupna-zmluva.docx", [], None)
    assert not _target_carries_pii(r"\\fileserver\vzory\zmluva.docx", [], None)


# FIXED 2026-09-17, daytime run. The marker is gone rather than flipped to xpass: a finding
# that has been fixed must become an ordinary regression test, or a later regression puts it
# back to "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
def test_tel_target_is_scrubbed():
    assert _target_carries_pii("tel:+421905123456", [], None)
    assert _target_carries_pii("tel:+421 905 123 456", [], None)


# ========================================= N1: the R3-A10 fix re-opened the homoglyph attack
CYR = {"S": "Ѕ", "K": "К", "A": "А", "B": "В", "E": "Е"}


def _cyr(s):
    return "".join(CYR.get(ch, ch) for ch in s)


def test_control_homoglyph_inside_a_latin_word_is_still_folded():
    """N1's control: the defence the R3-A10 fix was built to keep must still work."""
    assert normalize(_cyr("Kosice")).text == "Kosice"


# FIXED 2026-09-17, overnight run. The marker is gone rather than flipped to xpass: a finding that has been
# fixed must become an ordinary regression test, or a later regression puts it back to
# "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
@pytest.mark.parametrize("text,want", [
    ("Uhradte na ucet SK3112000000198742637541 do 30 dni.", "IBAN"),
    ("Vozidlo ECV BA123AB je vo vlastnictve.", "ECV"),
    ("Cislo pasu: EA123456 platny do roku 2030.", "CISLO_OP"),
])
def test_homoglyph_identifier_is_still_detected(text, want):
    mutated = " ".join(_cyr(t) if any(c.isdigit() for c in t) else t for t in text.split())
    assert want in {c.type for c in detect(text, [])}      # premise
    assert want in {c.type for c in detect(mutated, [])}


# =========================================================================== X: cross-cutting
@pytest.mark.xfail(reason="R4-X2: detect() is quadratic in the length of a single-case unit "
                          "whose lines are one capitalised token each — an all-caps column of "
                          "surnames, which is what a land-registry owner list looks like",
                   strict=True)
def test_detect_cost_is_not_quadratic_on_an_allcaps_name_column():
    names = ["NOVAK", "KOVACOVA", "HORVATH", "BALAZ", "TOTHOVA",
             "SZABO", "MOLNAR", "VARGA", "LUKACOVA", "BENKO"]

    def t(rows):
        s = "".join(names[i % 10] + "\n" for i in range(rows))
        detect(s, [])                       # warm
        t0 = time.perf_counter()
        detect(s, [])
        return time.perf_counter() - t0

    small, big = t(400), t(1600)            # x4 length
    assert big / max(small, 1e-6) < 8, (
        f"x4 input length cost x{big / max(small, 1e-6):.1f} time "
        f"({small * 1000:.0f} ms -> {big * 1000:.0f} ms)")


@pytest.mark.xfail(reason="R4-X1: the same letter-spaced page is REFUSED as a PDF and "
                          "silently under-redacted as a DOCX — two opposite verdicts, "
                          "neither of them 'redacted'", strict=True)
def test_docx_and_pdf_agree_on_a_letterspaced_page(tmp_path):
    lines = [SP("NOTARSKA") + "  " + SP("ZAPISNICA"),
             "Predavajuci:  " + SP("Jan") + "   " + SP("Novak"),
             "spisana dna 17. septembra 2026"]
    pdf = write_pdf(tmp_path, "x1", [(72, 90 + 24 * i, A(l)) for i, l in enumerate(lines)])
    pdf_refused = False
    try:
        redact_pdf(pdf, str(tmp_path / "x1_r.pdf"), ["Jan Novak"])
    except NoTextLayerError:
        pdf_refused = True
    out, lm = redact(tmp_path, "x1d",
                     [f'<w:p><w:r><w:t xml:space="preserve">{l}</w:t></w:r></w:p>'
                      for l in lines], ["Jan Novak"])
    docx_redacted = bool(lm.occurrences)
    # The two formats must AGREE. Today the PDF is refused and the DOCX is accepted with
    # nothing redacted, so the lawyer's answer depends on which copy they happened to open.
    assert pdf_refused is not True or docx_redacted is not False, (
        f"pdf REFUSED but the same content as DOCX was accepted with nothing redacted "
        f"(pdf_refused={pdf_refused}, docx_redacted={docx_redacted}, "
        f"docx still holds the name: {'J a n   N o v a k' in doc_text(out)})")


# FIXED 2026-09-17, daytime run. The marker is gone rather than flipped to xpass: a finding
# that has been fixed must become an ordinary regression test, or a later regression puts it
# back to "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
#
# The assertion is STRONGER than the one the round wrote. The original only failed on a raw
# XMLSyntaxError, which meant any other exception -- including the named refusal the finding
# asked for -- also failed the test, so it could never have gone green. What the contract
# actually promises is a NAMED refusal whose message says nothing was written, so that is what
# is pinned: the exception type, and the absence of the output file.
def test_malformed_rels_gives_a_clean_refusal(tmp_path):
    def patch(a, b):
        return patch_package(
            a, b, replace={"word/_rels/document.xml.rels": b"<Relationships><oops"})

    out = tmp_path / "x3_r.docx"
    with pytest.raises(UnreadableDocumentError):
        redact(tmp_path, "x3", [
            '<w:p><w:r><w:t xml:space="preserve">Predávajúci: Ján Novák.</w:t></w:r></w:p>'],
            KNOWN, patch=patch)
    assert not out.exists(), "a half-redacted file was left on disk"


def test_malformed_rels_deeper_in_the_package_also_refuses(tmp_path):
    """The same promise for a .rels part python-docx does not parse on open, which is the
    branch `_scrub_rel_targets` itself has to guard: it reaches that function as bytes and
    used to raise lxml's own exception out of the writer."""
    def patch(a, b):
        return patch_package(
            a, b, replace={"customXml/_rels/item1.xml.rels": b"<Relationships><oops"})

    out = tmp_path / "x3b_r.docx"
    with pytest.raises(UnreadableDocumentError):
        redact(tmp_path, "x3b", [
            '<w:p><w:r><w:t xml:space="preserve">Predávajúci: Ján Novák.</w:t></w:r></w:p>'],
            KNOWN, patch=patch)
    assert not out.exists(), "a half-redacted file was left on disk"
