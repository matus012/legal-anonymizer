"""Red-team round 5 — WORD TEMPLATE STRUCTURES, hand-authored as raw OOXML.

python-docx cannot emit most of the shapes attacked here (building-block glossaries,
``w:altChunk``, first-page/even-page header references, legacy ``w:ffData`` form fields,
``w:ptab``), which is exactly why four rounds of real bugs left them untouched. Every fixture
below is hand-authored XML zipped into a valid OPC package, or lxml surgery on the parts of
one.

Every test asserts THE PROPERTY THAT SHOULD HOLD. A test that asserted the bug would be green
today and turn red the day somebody fixed it, which is backwards. Unfixed findings carry
``xfail(strict=True)`` with the finding id, so the day a fix lands the suite says so by
XPASSing.

The tests that PASS are CONTROLS, and they come in two kinds:

  * **reachability controls** — the needle is provably present on the expected surface of the
    UNREDACTED fixture, so "not found in the output" can never be confused with a malformed
    fixture or a blind extractor;
  * **contrast controls** — the identical construct in the location the writer DOES cover is
    redacted correctly, so a finding is attributable to the missing traversal rather than to a
    detector that never fired.

Every package built here is re-opened through python-docx and every XML part re-parsed before
anything is asserted about it (``reopens``): a fixture Word would reject proves nothing, and
neither does an output it would reject.

Findings are documented in ``redteam/FINDINGS_ROUND5.md``.
"""
from __future__ import annotations

import os
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from lxml import etree

from eval.extract import extract
from eval.leak import Leak
from eval.leak_gate import split_found_in
from writer.docx_body import _scrub_field_instruction, redact_docx_collect
from writer.errors import EmbeddedSubDocumentError, UnreadableDocumentError

# --------------------------------------------------------------------------------- OOXML kit
NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:o="urn:schemas-microsoft-com:office:office" '
    'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" '
    'xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml" '
)

KNOWN = ["Ján Novák", "Mária Kováčová"]

REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CT_ML = "application/vnd.openxmlformats-officedocument.wordprocessingml"


def doc_with_body(fragments, path, sect_extra=""):
    """A real python-docx package whose body is replaced by hand-authored fragments."""
    d = Document()
    body = d.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    sect = body.find(qn("w:sectPr"))
    for frag in fragments:
        for child in list(parse_xml(f"<w:_wrap {NS}>{frag}</w:_wrap>")):
            sect.addprevious(child) if sect is not None else body.append(child)
    if sect_extra and sect is not None:
        for child in list(parse_xml(f"<w:_wrap {NS}>{sect_extra}</w:_wrap>")):
            sect.insert(0, child)
    d.save(str(path))
    return str(path)


def patch_package(src, dst, *, add=None, replace=None, content_types=None, doc_rels=None):
    """Add / replace parts of a saved package and re-zip it."""
    add, replace = add or {}, replace or {}
    with zipfile.ZipFile(src) as z:
        items = [(n, z.read(n)) for n in z.namelist()]
    out = []
    for name, data in items:
        data = replace.get(name, data)
        if name == "[Content_Types].xml" and content_types:
            data = data.replace(b"</Types>", ("".join(content_types) + "</Types>").encode())
        if name == "word/_rels/document.xml.rels" and doc_rels:
            data = data.replace(b"</Relationships>",
                                ("".join(doc_rels) + "</Relationships>").encode())
        out.append((name, data))
    have = {n for n, _ in out}
    out += [(n, d) for n, d in add.items() if n not in have]
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in out:
            z.writestr(name, data)
    return str(dst)


def override(partname, kind):
    return f'<Override PartName="{partname}" ContentType="{CT_ML}.{kind}+xml"/>'


def relationship(rid, kind, target):
    return f'<Relationship Id="{rid}" Type="{REL}/{kind}" Target="{target}"/>'


def reopens(path):
    """Word's own terms, as far as this repo can check them without Word: the package re-opens
    through python-docx (so every part it needs parses and every relationship it follows
    resolves) and every XML/rels part in it is well-formed."""
    assert Document(path).element.body is not None
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        assert {"[Content_Types].xml", "_rels/.rels", "word/document.xml",
                "word/_rels/document.xml.rels"} <= names
        for name in names:
            if name.endswith(".xml") or name.endswith(".rels"):
                etree.fromstring(z.read(name))
    return True


def part(path, name):
    with zipfile.ZipFile(path) as z:
        return z.read(name).decode("utf-8", "replace")


def surfaces_with(path, needle):
    return sorted(s for s, t in extract(path).by_surface.items() if needle in t)


def redact(tmp_path, name, fragments, known=KNOWN, patch=None, sect_extra=""):
    src = doc_with_body(fragments, Path(tmp_path) / f"{name}.docx", sect_extra=sect_extra)
    if patch is not None:
        src = patch(src, str(Path(tmp_path) / f"{name}_p.docx"))
    assert reopens(src), "fixture control: the SOURCE package must be well-formed"
    out = str(Path(tmp_path) / f"{name}_r.docx")
    lm = redact_docx_collect(src, out, list(known))
    assert reopens(out), "output control: the REDACTED package must still be well-formed"
    return out, lm


# =========================================================== R5-01  first / even page headers
HDR_TMPL = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<w:{{tag}} {NS}><w:p><w:r><w:t xml:space="preserve">{{txt}}</w:t></w:r>'
            '</w:p></w:{tag}>')

_HF_PARTS = {
    "word/header1.xml": ("hdr", "Advokat: Ján Novák, default"),
    "word/header2.xml": ("hdr", "Klient: Mária Kováčová, prva strana"),
    "word/header3.xml": ("hdr", "Klient: Mária Kováčová, parna strana"),
    "word/footer1.xml": ("ftr", "Ján Novák, default pata"),
    "word/footer2.xml": ("ftr", "Ján Novák, prva pata"),
}
_HF_SECT = ('<w:titlePg/>'
            '<w:headerReference w:type="default" r:id="rIdH1"/>'
            '<w:headerReference w:type="first" r:id="rIdH2"/>'
            '<w:headerReference w:type="even" r:id="rIdH3"/>'
            '<w:footerReference w:type="default" r:id="rIdF1"/>'
            '<w:footerReference w:type="first" r:id="rIdF2"/>')


def _patch_headers(src, dst):
    add = {name: HDR_TMPL.format(tag=tag, txt=txt).encode()
           for name, (tag, txt) in _HF_PARTS.items()}
    cts = [override(f"/{n}", "header" if tag == "hdr" else "footer")
           for n, (tag, _) in _HF_PARTS.items()]
    rels = [relationship(f"rIdH{i}", "header", f"header{i}.xml") for i in (1, 2, 3)]
    rels += [relationship(f"rIdF{i}", "footer", f"footer{i}.xml") for i in (1, 2)]
    return patch_package(src, dst, add=add, content_types=cts, doc_rels=rels)


def _headers_fixture(tmp_path):
    return redact(tmp_path, "r5_01", ['<w:p><w:r><w:t>Telo dokumentu.</w:t></w:r></w:p>'],
                  patch=_patch_headers, sect_extra=_HF_SECT)


def test_control_r5_01_default_header_and_footer_are_redacted(tmp_path):
    """CONTRAST CONTROL for R5-01. The DEFAULT header and footer go through the writer, so a
    miss on the first/even parts is the traversal, not the detector."""
    out, lm = _headers_fixture(tmp_path)
    assert "[MENO_1]" in part(out, "word/header1.xml")
    assert "Ján Novák" not in part(out, "word/header1.xml")
    assert "[MENO_1]" in part(out, "word/footer1.xml")
    assert ("header", "Ján Novák") in lm.occurrences["[MENO_1]"]


def test_control_r5_01_needles_are_reachable_in_the_source(tmp_path):
    """REACHABILITY CONTROL for R5-01: every needle is on the `header`/`footer` surface of the
    UNREDACTED package, so "absent from the output" can only mean "removed"."""
    src = doc_with_body(['<w:p><w:r><w:t>Telo dokumentu.</w:t></w:r></w:p>'],
                        tmp_path / "r5_01_src.docx", sect_extra=_HF_SECT)
    src = _patch_headers(src, str(tmp_path / "r5_01_src_p.docx"))
    assert reopens(src)
    assert surfaces_with(src, "Mária Kováčová") == ["header"]
    assert surfaces_with(src, "Ján Novák") == ["footer", "header"]


# FIXED 2026-09-17, daytime run. The marker is gone rather than flipped to xpass: a finding
# that has been fixed must become an ordinary regression test, or a later regression puts it
# back to "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
def test_r5_01_first_and_even_page_headers_and_footers_are_redacted(tmp_path):
    out, _ = _headers_fixture(tmp_path)
    survivors = []
    for name, (_, txt) in _HF_PARTS.items():
        body = part(out, name)
        for needle in ("Ján Novák", "Mária Kováčová"):
            if needle in body:
                survivors.append((name, needle))
    assert survivors == [], f"PII survived in non-default header/footer parts: {survivors}"


# ================================================================= R5-02  w:altChunk
_MHT = (
    "MIME-Version: 1.0\r\nContent-Type: text/html; charset=\"utf-8\"\r\n\r\n"
    "<html><body><p>Kupujuci: Mária Kováčová, rodné číslo 855612/7788, PSC 04001, "
    "kod banky 1100.</p></body></html>\r\n"
).encode("utf-8")


def _patch_altchunk(src, dst):
    return patch_package(
        src, dst,
        add={"word/afchunk.mht": _MHT},
        content_types=['<Default Extension="mht" ContentType="message/rfc822"/>'],
        doc_rels=[relationship("rIdAC", "aFChunk", "afchunk.mht")],
    )


_AC_BODY = ['<w:p><w:r><w:t>Predavajuci: Ján Novák.</w:t></w:r></w:p>',
            '<w:altChunk r:id="rIdAC"/>']


def _altchunk_source(tmp_path, name="r5_02"):
    """The SOURCE package: an ordinary body paragraph plus an altChunk pointing at an .mht.

    The fixture is built and checked here rather than through ``redact()`` because the writer
    now REFUSES this document (see below), so there is no output to assert about. Everything
    ``redact()`` guarantees about a source is still asserted: it re-opens through python-docx
    and every XML/rels part in it parses.
    """
    src = doc_with_body(_AC_BODY, Path(tmp_path) / f"{name}.docx")
    src = _patch_altchunk(src, str(Path(tmp_path) / f"{name}_p.docx"))
    assert reopens(src), "fixture control: the SOURCE package must be well-formed"
    return src


# FIXED 2026-09-17, second run. THE CONTRACT CHANGED, so these two controls changed with it.
# R5-02 was written expecting the chunk to be REDACTED, and both controls were phrased against
# an OUTPUT file. The fix is a REFUSAL (writer/errors.py::EmbeddedSubDocumentError): an altChunk
# is not WordprocessingML, this writer cannot rewrite it, and silently deleting a merged-in
# annex is document damage the reviewer cannot see. There is therefore no output package to
# inspect, and a control that asserts one would assert the bug back into existence.
#
# Each control keeps its ORIGINAL PURPOSE, restated against what now exists:
#   * the contrast control still proves the refusal is attributable to the chunk and not to the
#     body -- by showing the same body, with the altChunk removed, is redacted normally;
#   * the reachability control still proves the needle is visible to the extractor, and now
#     pins the surface it must be visible ON, which is the whole of the R5-02b half.
def test_control_r5_02_the_same_body_without_a_chunk_is_redacted(tmp_path):
    """CONTRAST CONTROL for R5-02: the identical body paragraph in a document with NO altChunk
    goes through the writer untouched by the refusal, so the refusal is about the chunk."""
    out, lm = redact(tmp_path, "r5_02_ctl", _AC_BODY[:1])
    assert surfaces_with(out, "Ján Novák") == []
    assert lm.occurrences["[MENO_1]"] == [("body", "Ján Novák")]


def test_control_r5_02_the_chunk_part_is_reachable_and_is_not_an_opaque_surface(tmp_path):
    """REACHABILITY CONTROL for R5-02: the extractor sees the chunk's text, and sees it on
    `alt_chunk_parts` -- a TEXT surface -- not on `binary_parts`, which the leak gate discounts
    short needles on. Before the R5-02 extractor fix this read `["binary_parts"]`."""
    src = _altchunk_source(tmp_path, "r5_02_reach")
    with zipfile.ZipFile(src) as z:
        assert "word/afchunk.mht" in z.namelist()
    assert surfaces_with(src, "Mária Kováčová") == ["alt_chunk_parts"]


# FIXED 2026-09-17, second run — as a REFUSAL, which is what the finding itself recommended and
# is the shape test_r5_10 already contracts for ("redacted or refused by name"). The marker is
# gone rather than flipped to xpass, for the reason the other fixed findings in this file give.
def test_r5_02_altchunk_content_is_redacted_or_the_document_is_refused_by_name(tmp_path):
    src = _altchunk_source(tmp_path, "r5_02_fix")
    out = str(Path(tmp_path) / "r5_02_fix_r.docx")
    try:
        redact_docx_collect(src, out, list(KNOWN))
    except EmbeddedSubDocumentError as exc:
        # The message promises nothing was written. That has to be TRUE on disk, and it has to
        # say what an altChunk is in terms the reviewer can act on.
        assert not os.path.exists(out), "refused, but a partial output was left on disk"
        assert "word/afchunk.mht" in str(exc)
        assert "embedded sub-document" in str(exc)
        assert isinstance(exc, UnreadableDocumentError), "every caller that refuses an " \
            "unreadable document must refuse this one unchanged"
        return
    chunk = part(out, "word/afchunk.mht")
    survivors = [n for n in ("Mária Kováčová", "855612/7788", "04001") if n in chunk]
    assert survivors == [], f"embedded altChunk document shipped unredacted: {survivors}"


def test_r5_02_an_ordinary_document_is_not_refused_as_an_embedded_sub_document(tmp_path):
    """FALSE-REFUSAL CONTROL. A guard that fires on ordinary work is a guard the office turns
    off. Nothing without an aFChunk relationship may reach the refusal."""
    out, _ = redact(tmp_path, "r5_02_false", _AC_BODY[:1])
    assert os.path.exists(out)


def test_r5_02_a_nested_docx_altchunk_is_readable_and_refused(tmp_path):
    """The OTHER altChunk flavour, and a worse one for the gate: a whole nested .docx. Its
    document.xml is DEFLATED inside the inner ZIP inside the outer ZIP, so no decoding of the
    outer bytes could see a word of it — on `binary_parts` it was not merely discounted, it was
    INVISIBLE at any needle length. It must now read as text, and be refused."""
    inner = doc_with_body([], tmp_path / "inner.docx")
    d = Document(inner)
    d.add_paragraph("Kupujuci: Maria Kovacova, PSC 04001, kod banky 1100.")
    d.save(inner)
    src = doc_with_body(['<w:p><w:r><w:t>Telo.</w:t></w:r></w:p>', '<w:altChunk r:id="rIdAC"/>'],
                        tmp_path / "outer.docx")
    src = patch_package(
        src, str(tmp_path / "outer_p.docx"),
        add={"word/afchunk.docx": Path(inner).read_bytes()},
        content_types=['<Default Extension="docx" ContentType="application/vnd.openxmlformats-'
                       'officedocument.wordprocessingml.document"/>'],
        doc_rels=[relationship("rIdAC", "aFChunk", "afchunk.docx")],
    )
    assert reopens(src)
    assert surfaces_with(src, "04001") == ["alt_chunk_parts"]
    assert surfaces_with(src, "Maria Kovacova") == ["alt_chunk_parts"]
    out = str(tmp_path / "outer_r.docx")
    with pytest.raises(EmbeddedSubDocumentError):
        redact_docx_collect(src, out, ["Maria Kovacova"])
    assert not os.path.exists(out)


# FIXED 2026-09-17, second run (R5-02b, the extractor half). Graded on the SOURCE package: the
# writer now refuses to emit a package containing an altChunk, so the source is the only place
# an altChunk can be graded -- and grading it is exactly the point, because it is what would
# catch a regression that dropped the refusal.
def test_r5_02_the_leak_gate_grades_a_short_needle_in_an_altchunk(tmp_path):
    src = _altchunk_source(tmp_path, "r5_02_gate")
    res = extract(src)
    ungraded = []
    for needle, ptype in (("04001", "PSC"), ("1100", "KOD_BANKY")):
        found = tuple(s for s, t in res.by_surface.items() if needle in t)
        assert found, f"fixture control: {needle!r} must be present somewhere"
        counted, _aside = split_found_in(Leak(needle, ptype, found, ("body",)), res.by_surface)
        if not counted:
            ungraded.append((needle, found))
    assert ungraded == [], f"leak gate scores these CLEAN: {ungraded}"


# ================================================================= R5-03  glossary document
_GLOSSARY = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    f'<w:glossaryDocument {NS}><w:docParts><w:docPart><w:docPartPr>'
    '<w:name w:val="Titulna strana kancelarie"/>'
    '<w:category><w:name w:val="General"/><w:gallery w:val="coverPg"/></w:category>'
    '<w:types><w:type w:val="bbPlcHdr"/></w:types></w:docPartPr><w:docPartBody>'
    '<w:p><w:ins w:id="1" w:author="JUDr. Ján Novák" w:date="2026-01-01T00:00:00Z">'
    '<w:r><w:t xml:space="preserve">Klient: Mária Kováčová, rodné číslo 855612/7788</w:t>'
    '</w:r></w:ins>'
    '<w:fldSimple w:instr=" HYPERLINK &quot;mailto:maria.kovacova@gmail.com&quot; ">'
    '<w:r><w:t xml:space="preserve">kontakt</w:t></w:r></w:fldSimple></w:p>'
    '</w:docPartBody></w:docPart></w:docParts></w:glossaryDocument>'
)
# The shape every Word template ships with: a cover page assembled from a building block.
_COVER_SDT = (
    '<w:sdt><w:sdtPr><w:id w:val="12345"/>'
    '<w:docPartObj><w:docPartGallery w:val="Cover Pages"/><w:docPartUnique/></w:docPartObj>'
    '</w:sdtPr><w:sdtContent>'
    '<w:p><w:r><w:t xml:space="preserve">Klient: Mária Kováčová</w:t></w:r></w:p>'
    '</w:sdtContent></w:sdt>'
)


def _patch_glossary(src, dst):
    return patch_package(
        src, dst,
        add={"word/glossary/document.xml": _GLOSSARY.encode()},
        content_types=[f'<Override PartName="/word/glossary/document.xml" '
                       f'ContentType="{CT_ML}.document.glossary+xml"/>'],
        doc_rels=[relationship("rIdGloss", "glossaryDocument", "glossary/document.xml")],
    )


def _glossary_fixture(tmp_path):
    return redact(tmp_path, "r5_03", [_COVER_SDT], patch=_patch_glossary)


def test_control_r5_03_the_cover_page_sdt_in_the_body_is_redacted(tmp_path):
    """CONTRAST CONTROL for R5-03: a `w:sdt` / `w:docPartObj` cover page IN THE BODY is
    redacted, so the glossary miss is about the PART, not about content controls.

    AMENDED 2026-09-17 with the R5-03 fix. The occurrence list was pinned to exactly
    ``[("body", ...)]``, which was only true while the glossary went unvisited; the same name is
    now redacted in the glossary too and is recorded there under its own location. The body
    occurrence is still pinned FIRST (traversal order is body before the glossary), which is
    what this control is for."""
    out, lm = _glossary_fixture(tmp_path)
    assert "Mária Kováčová" not in part(out, "word/document.xml")
    assert lm.occurrences["[MENO_1]"] == [("body", "Mária Kováčová"),
                                          ("glossary", "Mária Kováčová")]


def test_control_r5_03_the_glossary_part_is_reachable_by_the_extractor(tmp_path):
    """REACHABILITY CONTROL for R5-03: `other_xml_parts` is a TEXT surface, so a leak here is
    one the gate would grade — it was simply never scrubbed.

    AMENDED 2026-09-17 with the R5-03 fix. As written this measured the REDACTED output, so it
    asserted the leak was still there — it was a control that could only hold while the finding
    was open. A reachability control measures the UNREDACTED SOURCE (as R5-01's does): the
    needle is provably on a graded surface of the input, so "absent from the output" can only
    mean "removed"."""
    src = doc_with_body([_COVER_SDT], tmp_path / "r5_03_src.docx")
    src = _patch_glossary(src, str(tmp_path / "r5_03_src_p.docx"))
    assert reopens(src)
    assert surfaces_with(src, "855612/7788") == ["other_xml_parts"]


# FIXED 2026-09-17, daytime run. The marker is gone rather than flipped to xpass: a finding
# that has been fixed must become an ordinary regression test, or a later regression puts it
# back to "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
def test_r5_03_the_building_block_glossary_is_redacted(tmp_path):
    out, _ = _glossary_fixture(tmp_path)
    gloss = part(out, "word/glossary/document.xml")
    survivors = [n for n in ("Mária Kováčová", "855612/7788", "JUDr. Ján Novák",
                             "maria.kovacova@gmail.com") if n in gloss]
    assert survivors == [], f"building-block glossary shipped unredacted: {survivors}"


# ============================================== R5-04  field codes in footnote / endnote parts
_FOOTNOTES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    f'<w:footnotes {NS}>'
    '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
    '<w:footnote w:id="1"><w:p>'
    '<w:r><w:t xml:space="preserve">Pozri zmluvu s Ján Novák, </w:t></w:r>'
    '<w:fldSimple w:instr=" HYPERLINK &quot;mailto:maria.kovacova@gmail.com&quot; ">'
    '<w:r><w:t xml:space="preserve">kontakt</w:t></w:r></w:fldSimple>'
    '<w:r><w:instrText xml:space="preserve"> HYPERLINK '
    '"https://dms.local/klienti/Jan-Novak/zmluva.pdf" </w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
    '<w:r><w:t xml:space="preserve">spis</w:t></w:r>'
    '</w:p></w:footnote></w:footnotes>'
)
_ENDNOTES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    f'<w:endnotes {NS}><w:endnote w:id="1"><w:p>'
    '<w:r><w:instrText xml:space="preserve"> HYPERLINK '
    '"mailto:maria.kovacova@gmail.com" </w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
    '<w:r><w:t xml:space="preserve">kontakt</w:t></w:r></w:p></w:endnote></w:endnotes>'
)
# The brief's shape: the reference lives in a TABLE CELL, the note text in a separate OPC part.
_TABLE_WITH_NOTE_REF = (
    '<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
    '<w:tr><w:tc><w:tcPr><w:tcW w:w="4000" w:type="dxa"/></w:tcPr>'
    '<w:p><w:r><w:t xml:space="preserve">Kupujuci</w:t></w:r>'
    '<w:r><w:footnoteReference w:id="1"/></w:r>'
    '<w:r><w:endnoteReference w:id="1"/></w:r></w:p></w:tc></w:tr></w:tbl>'
)


def _patch_notes(src, dst):
    return patch_package(
        src, dst,
        add={"word/footnotes.xml": _FOOTNOTES.encode(),
             "word/endnotes.xml": _ENDNOTES.encode()},
        content_types=[override("/word/footnotes.xml", "footnotes"),
                       override("/word/endnotes.xml", "endnotes")],
        doc_rels=[relationship("rIdFn", "footnotes", "footnotes.xml"),
                  relationship("rIdEn", "endnotes", "endnotes.xml")],
    )


def _notes_fixture(tmp_path):
    return redact(tmp_path, "r5_04",
                  [_TABLE_WITH_NOTE_REF, '<w:p><w:r><w:t>Telo.</w:t></w:r></w:p>'],
                  patch=_patch_notes)


def test_control_r5_04_note_body_text_is_redacted(tmp_path):
    """CONTRAST CONTROL for R5-04: ordinary TEXT in the footnote part IS redacted, so
    _redact_notes_part reaches the part; only the field pass does not run on it."""
    out, lm = _notes_fixture(tmp_path)
    assert "Ján Novák" not in part(out, "word/footnotes.xml")
    assert ("footnote", "Ján Novák") in lm.occurrences["[MENO_1]"]


def test_control_r5_04_the_same_field_codes_in_the_body_are_scrubbed(tmp_path):
    """CONTRAST CONTROL for R5-04: both of Word's field spellings are scrubbed in the BODY.
    The footnote miss is therefore the missing _scrub_field_codes call, not the scrub."""
    out, _ = redact(tmp_path, "r5_04_body", [
        '<w:p><w:fldSimple w:instr=" HYPERLINK &quot;mailto:maria.kovacova@gmail.com&quot; ">'
        '<w:r><w:t xml:space="preserve">kontakt</w:t></w:r></w:fldSimple></w:p>',
        '<w:p><w:r><w:instrText xml:space="preserve"> HYPERLINK '
        '"https://dms.local/klienti/Jan-Novak/zmluva.pdf" </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r></w:p>'])
    body = part(out, "word/document.xml")
    assert "maria.kovacova@gmail.com" not in body
    assert "Jan-Novak" not in body


# FIXED 2026-09-17, daytime run. The marker is gone rather than flipped to xpass: a finding
# that has been fixed must become an ordinary regression test, or a later regression puts it
# back to "xfail" -- the state this file calls normal -- and nobody notices the fix was undone.
def test_r5_04_field_instructions_in_note_parts_are_scrubbed(tmp_path):
    out, _ = _notes_fixture(tmp_path)
    survivors = []
    for name in ("word/footnotes.xml", "word/endnotes.xml"):
        body = part(out, name)
        for needle in ("maria.kovacova@gmail.com", "Jan-Novak"):
            if needle in body:
                survivors.append((name, needle))
    assert survivors == [], f"field-code destinations survived in note parts: {survivors}"


# ====================================================== R5-05  docProps/app.xml beyond two tags
_APP = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"'
    ' xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
    '<Template>Normal.dotm</Template><Company>Advokatska kancelaria</Company>'
    '<Manager>Ján Novák</Manager>'
    '<HeadingPairs><vt:vector size="2" baseType="variant">'
    '<vt:variant><vt:lpstr>Nadpisy</vt:lpstr></vt:variant>'
    '<vt:variant><vt:i4>2</vt:i4></vt:variant></vt:vector></HeadingPairs>'
    '<TitlesOfParts><vt:vector size="2" baseType="lpstr">'
    '<vt:lpstr>Kúpna zmluva – Mária Kováčová</vt:lpstr>'
    '<vt:lpstr>Splnomocnenie pre Petra Horvátha</vt:lpstr>'
    '</vt:vector></TitlesOfParts>'
    '<HyperlinkBase>\\\\fileserver\\klienti\\kovacova.maria</HyperlinkBase>'
    '</Properties>'
)


def _app_fixture(tmp_path, blob, name):
    return redact(tmp_path, name, ['<w:p><w:r><w:t>Telo.</w:t></w:r></w:p>'],
                  patch=lambda s, d: patch_package(s, d, replace={"docProps/app.xml": blob}))


def test_control_r5_05_company_and_manager_are_blanked(tmp_path):
    """CONTRAST CONTROL for R5-05: the two tags _APP_XML_PII_TAGS names ARE blanked."""
    out, _ = _app_fixture(tmp_path, _APP.encode(), "r5_05_ctl")
    app = part(out, "docProps/app.xml")
    assert "<Company></Company>" in app and "<Manager></Manager>" in app
    assert "Ján Novák" not in app


def test_control_r5_05_app_xml_is_a_named_text_surface(tmp_path):
    """REACHABILITY CONTROL for R5-05: `app_xml` is a NAMED, non-opaque surface, so a leak here
    is graded by the leak gate with no attributability discount at all."""
    out, _ = _app_fixture(tmp_path, _APP.encode(), "r5_05_reach")
    res = extract(out)
    found = tuple(s for s, t in res.by_surface.items() if "Mária Kováčová" in t)
    assert found == ("app_xml",)
    counted, aside = split_found_in(Leak("Mária Kováčová", "MENO", found, ("metadata_app",)),
                                    res.by_surface)
    assert counted == ("app_xml",) and aside == ()


@pytest.mark.xfail(strict=True, reason="R5-05: _APP_XML_PII_TAGS is (Company, Manager) only; "
                                       "Word writes every heading into TitlesOfParts and the "
                                       "office share path into HyperlinkBase")
def test_r5_05_app_xml_free_text_carrying_pii_is_blanked(tmp_path):
    out, _ = _app_fixture(tmp_path, _APP.encode(), "r5_05")
    app = part(out, "docProps/app.xml")
    survivors = [n for n in ("Mária Kováčová", "Petra Horvátha", "kovacova.maria") if n in app]
    assert survivors == [], f"docProps/app.xml shipped PII: {survivors}"


# ======================================================== R5-06  app.xml scrub is prefix-blind
_APP_PREFIXED = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<ep:Properties '
    'xmlns:ep="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
    'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
    '<ep:Company>Ján Novák s.r.o.</ep:Company>'
    '<ep:Manager>Mária Kováčová</ep:Manager>'
    '</ep:Properties>'
)


def test_control_r5_06_the_prefixed_part_is_the_same_document(tmp_path):
    """FIXTURE CONTROL for R5-06: a prefixed serialisation is the SAME XML infoset — the
    elements resolve to the same expanded names as the unprefixed one Word writes."""
    ep = "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}"
    a = etree.fromstring(_APP.encode())
    b = etree.fromstring(_APP_PREFIXED.encode())
    assert a.find(ep + "Company") is not None
    assert b.find(ep + "Company") is not None
    assert b.find(ep + "Manager").text == "Mária Kováčová"


@pytest.mark.xfail(strict=True, reason="R5-06: _scrub_metadata rewrites app.xml with a regex "
                                       "over the raw bytes ('<Company>.+?</Company>'), which "
                                       "is blind to any namespace prefix")
def test_r5_06_app_xml_scrub_is_namespace_prefix_independent(tmp_path):
    out, _ = _app_fixture(tmp_path, _APP_PREFIXED.encode(), "r5_06")
    app = part(out, "docProps/app.xml")
    survivors = [n for n in ("Ján Novák", "Mária Kováčová") if n in app]
    assert survivors == [], f"prefixed app.xml was not scrubbed: {survivors}"


# =========================================== R5-07  attribute-borne PII nothing sweeps
_ATTRS = (
    # DrawingML picture alt text — the scanned signature block of a Slovak filing
    '<w:p><w:r><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
    '<wp:extent cx="1" cy="1"/>'
    '<wp:docPr id="7" name="Podpis" descr="Podpis klienta: Ján Novák" title="Ján Novák"/>'
    '<a:graphic><a:graphicData '
    'uri="http://schemas.openxmlformats.org/drawingml/2006/picture"/></a:graphic>'
    '</wp:inline></w:drawing></w:r></w:p>'
    # legacy VML alt text, with no DrawingML twin at all
    '<w:p><w:r><w:pict><v:shape id="s1" alt="Peciatka: Mária Kováčová" '
    'title="Mária Kováčová" style="width:10pt;height:10pt"/></w:pict></w:r></w:p>'
    # a bookmark named after the party
    '<w:p><w:bookmarkStart w:id="3" w:name="Adresa_Maria_Kovacova"/>'
    '<w:r><w:t xml:space="preserve">Adresa strany.</w:t></w:r>'
    '<w:bookmarkEnd w:id="3"/></w:p>'
    # SDT-heavy template: showingPlcHdr means the PAGE shows the placeholder while the stored
    # selection and the whole client list sit in w:sdtPr attributes
    '<w:sdt><w:sdtPr><w:alias w:val="Klient: Ján Novák"/><w:tag w:val="klient_jan_novak"/>'
    '<w:id w:val="9"/><w:showingPlcHdr/>'
    '<w:placeholder><w:docPart w:val="DefaultPlaceholder_1081868574"/></w:placeholder>'
    '<w:dropDownList w:lastValue="Ján Novák">'
    '<w:listItem w:displayText="Ján Novák" w:value="Ján Novák"/>'
    '<w:listItem w:displayText="Mária Kováčová" w:value="Mária Kováčová"/>'
    '<w:listItem w:displayText="Peter Horváth" w:value="Peter Horváth"/>'
    '</w:dropDownList></w:sdtPr><w:sdtContent>'
    '<w:p><w:r><w:t xml:space="preserve">Kliknite sem a zadajte meno.</w:t></w:r></w:p>'
    '</w:sdtContent></w:sdt>'
    # legacy Developer-tab form field: the default value and the status bar text
    '<w:p><w:r><w:fldChar w:fldCharType="begin"><w:ffData>'
    '<w:name w:val="Klient"/><w:enabled/><w:calcOnExit w:val="0"/>'
    '<w:statusText w:type="text" w:val="Meno klienta: Ján Novák"/>'
    '<w:textInput><w:default w:val="Mária Kováčová"/></w:textInput>'
    '</w:ffData></w:fldChar></w:r>'
    '<w:r><w:instrText xml:space="preserve"> FORMTEXT </w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
    '<w:r><w:t xml:space="preserve">Mária Kováčová</w:t></w:r>'
    '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
)

_ATTR_NEEDLES = ("Podpis klienta: Ján Novák", "Peciatka: Mária Kováčová",
                 "Adresa_Maria_Kovacova", "Klient: Ján Novák", "klient_jan_novak",
                 "Peter Horváth", "Meno klienta: Ján Novák")


def test_control_r5_07_the_form_field_result_run_on_the_page_is_redacted(tmp_path):
    """CONTRAST CONTROL for R5-07: the run the reader actually sees IS redacted, which is what
    makes the attribute survival invisible to a reviewer."""
    out, lm = redact(tmp_path, "r5_07_ctl", [_ATTRS])
    assert "<w:t xml:space=\"preserve\">Mária Kováčová</w:t>" not in part(out, "word/document.xml")
    assert lm.occurrences["[MENO_1]"] == [("body", "Mária Kováčová")]


def test_control_r5_07_a_short_attribute_needle_is_set_aside_by_the_gate(tmp_path):
    """REACHABILITY CONTROL for R5-07: `xml_attributes` is an OPAQUE surface, so the gate only
    counts a needle of >= 7 characters. A first name or a 4-digit code here is CLEAN."""
    out, _ = redact(tmp_path, "r5_07_short",
                    ['<w:p><w:r><w:pict><v:shape id="s2" alt="Klient Ján" '
                     'style="width:10pt;height:10pt"/></w:pict></w:r></w:p>'])
    res = extract(out)
    found = tuple(s for s, t in res.by_surface.items() if "Ján" in t)
    assert "xml_attributes" in found
    counted, aside = split_found_in(Leak("Ján", "MENO", ("xml_attributes",), ("body",)),
                                    res.by_surface)
    assert counted == () and aside == ("xml_attributes",)


@pytest.mark.xfail(strict=True, reason="R5-07: the writer's only attribute sweep is "
                                       "w:author/w:initials; alt text, bookmark names, sdtPr "
                                       "aliases, dropdown lists and w:ffData defaults survive")
def test_r5_07_attribute_borne_pii_in_document_xml_is_removed(tmp_path):
    out, _ = redact(tmp_path, "r5_07", [_ATTRS])
    body = part(out, "word/document.xml")
    survivors = [n for n in _ATTR_NEEDLES if n in body]
    assert survivors == [], f"attribute values shipped PII: {survivors}"


# ============================================ R5-08  the quoted-bookmark hole in the R4 fix
_TOC = (
    '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
    '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="separate"/></w:r></w:p>'
    '<w:p><w:hyperlink w:anchor="_Toc53871234">'
    '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
    '<w:r><w:instrText xml:space="preserve"> HYPERLINK \\l "_Toc53871234" </w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
    '<w:r><w:t xml:space="preserve">Clanok I. Predmet zmluvy</w:t></w:r>'
    '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:hyperlink></w:p>'
    '<w:p><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
)


def test_control_r5_08_the_bare_bookmark_spelling_is_exempt(tmp_path):
    """CONTRAST CONTROL for R5-08: R4's fix works for the BARE spelling and for PAGEREF, and a
    real slov-lex citation still survives. The hole is the quoted argument only."""
    assert _scrub_field_instruction(' HYPERLINK \\l _Toc53871234 ', [], None) is None
    assert _scrub_field_instruction(' PAGEREF _Toc53871234 \\h ', [], None) is None
    assert _scrub_field_instruction(
        ' HYPERLINK "https://www.slov-lex.sk/pravne-predpisy/SK/ZZ/2016/18/20180101" ',
        [], None) is None


# FIXED 2026-09-17, daytime run, and this one was the ORCHESTRATOR'S OWN spec error rather
# than the implementing round's: the exemption was specified as "a BARE (unquoted) argument
# beginning with _", which is exactly half of how Word writes a bookmark name. REF and PAGEREF
# write it bare and are not examined at all; HYPERLINK \l writes it QUOTED and is. The value is
# what carries a bookmark name, so the value is now what is tested -- while the BACKSLASH test
# stays on the bare token, because a quoted argument starting with a backslash is a UNC path
# and is a destination this must examine. The two exemptions are not symmetric.
#
# The marker is gone rather than flipped to xpass: a finding that has been fixed must become an
# ordinary regression test, or a later regression puts it back to "xfail" -- the state this file
# calls normal -- and nobody notices the fix was undone.
def test_r5_08_a_quoted_word_bookmark_name_is_never_treated_as_a_destination(tmp_path):
    damaged = [instr for instr in (' HYPERLINK \\l "_Toc53871234" ',
                                   ' HYPERLINK \\l "_Ref53871234" ',
                                   ' HYPERLINK \\l "_Toc5387123456" ')
               if _scrub_field_instruction(instr, [], None) is not None]
    out, _ = redact(tmp_path, "r5_08", [_TOC])
    body = part(out, "word/document.xml")
    if "_Toc53871234" not in body:
        damaged.append("table of contents entry in the document")
    assert damaged == [], f"ordinary Word cross-references destroyed: {damaged}"


# ======================= R5-09  _TEXT_BEARING disagrees with python-docx's own run.text xpath
def test_control_r5_09_python_docx_run_text_is_what_the_offsets_are_built_from(tmp_path):
    """FIXTURE CONTROL for R5-09. _TEXT_BEARING must be exactly the set of children that
    CONTRIBUTE to Run.text. Measured against the installed python-docx: w:ptab contributes a
    tab and is MISSING from the set; w:softHyphen contributes nothing and is IN it."""
    def rtext(frag):
        return parse_xml(f'<w:r {NS}>{frag}</w:r>').text
    assert rtext('<w:t>A</w:t><w:ptab w:relativeTo="margin" w:alignment="left" '
                 'w:leader="none"/><w:t>B</w:t>') == "A\tB"
    assert rtext('<w:t>A</w:t><w:softHyphen/><w:t>B</w:t>') == "AB"
    assert rtext('<w:t>A</w:t><w:noBreakHyphen/><w:t>B</w:t>') == "A-B"


def test_control_r5_09_an_untouched_run_keeps_its_inner_content(tmp_path):
    """CONTRAST CONTROL for R5-09: a run with no redaction in it is left byte-identical, so the
    damage below is the rebuild path and nothing else."""
    out, _ = redact(tmp_path, "r5_09_ctl",
                    ['<w:p><w:r><w:t xml:space="preserve">Clanok I.</w:t>'
                     '<w:ptab w:relativeTo="margin" w:alignment="right" w:leader="dot"/>'
                     '<w:t xml:space="preserve">strana 1</w:t></w:r></w:p>'])
    body = part(out, "word/document.xml")
    assert body.count("w:ptab") == 1
    assert "<w:tab/>" not in body


@pytest.mark.xfail(strict=True, reason="R5-09: w:ptab is not in _TEXT_BEARING so it is cloned "
                                       "into every fragment AND re-emitted as w:tab; "
                                       "w:softHyphen is in it but renders as nothing, so it is "
                                       "deleted; w:noBreakHyphen is downgraded to a literal '-'")
def test_r5_09_rebuilding_a_run_preserves_its_inner_content_elements(tmp_path):
    faults = []

    out, _ = redact(tmp_path, "r5_09_ptab",
                    ['<w:p><w:r><w:t xml:space="preserve">Ján Novák</w:t>'
                     '<w:ptab w:relativeTo="margin" w:alignment="right" w:leader="dot"/>'
                     '<w:t xml:space="preserve">strana 1</w:t></w:r></w:p>'])
    body = part(out, "word/document.xml")
    if body.count("w:ptab") != 1:
        faults.append(f"w:ptab duplicated: {body.count('w:ptab')} copies")
    if "<w:tab/>" in body:
        faults.append("w:ptab additionally re-emitted as a plain w:tab")

    out, _ = redact(tmp_path, "r5_09_shy",
                    ['<w:p><w:r><w:t xml:space="preserve">Ján Novák nadobú</w:t>'
                     '<w:softHyphen/><w:t xml:space="preserve">da nehnutelnost</w:t></w:r></w:p>'])
    if "<w:softHyphen/>" not in part(out, "word/document.xml"):
        faults.append("w:softHyphen deleted from the rebuilt run")

    out, _ = redact(tmp_path, "r5_09_nbh",
                    ['<w:p><w:r><w:t xml:space="preserve">Ján Novák, c. zmluvy 12</w:t>'
                     '<w:noBreakHyphen/><w:t xml:space="preserve">2026</w:t></w:r></w:p>'])
    if "<w:noBreakHyphen/>" not in part(out, "word/document.xml"):
        faults.append("w:noBreakHyphen downgraded to a literal hyphen-minus")

    assert faults == [], f"run rebuild damaged inner content: {faults}"


# ================================ R5-10  a non-UTF-8 metadata part escapes as a codec error
_APP_CP1250 = (
    '<?xml version="1.0" encoding="windows-1250" standalone="yes"?>'
    '<Properties '
    'xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
    '<Company>Ján Novák</Company></Properties>'
).encode("cp1250")


def test_control_r5_10_the_cp1250_part_is_well_formed_xml(tmp_path):
    """FIXTURE CONTROL for R5-10: the part is valid XML — lxml honours its declaration and
    reads the Slovak text back. It is not malformed; it is simply not UTF-8."""
    root = etree.fromstring(_APP_CP1250)
    ep = "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}"
    assert root.find(ep + "Company").text == "Ján Novák"


@pytest.mark.xfail(strict=True, reason="R5-10: _scrub_metadata does part._blob.decode('utf-8') "
                                       "unguarded, so a legacy cp1250 docProps/app.xml leaves "
                                       "the writer as UnicodeDecodeError instead of the "
                                       "contracted UnreadableDocumentError")
def test_r5_10_a_non_utf8_metadata_part_is_redacted_or_refused_by_name(tmp_path):
    src = doc_with_body(['<w:p><w:r><w:t>Telo.</w:t></w:r></w:p>'], tmp_path / "r5_10.docx")
    src = patch_package(src, str(tmp_path / "r5_10_p.docx"),
                        replace={"docProps/app.xml": _APP_CP1250})
    out = str(tmp_path / "r5_10_r.docx")
    try:
        redact_docx_collect(src, out, list(KNOWN))
    except UnreadableDocumentError:
        pass  # the contracted refusal — acceptable
    except Exception as exc:  # noqa: BLE001 — that is the point of the test
        assert not os.path.exists(out), "and a partial output was left on disk"
        pytest.fail(f"writer raised {type(exc).__name__} instead of UnreadableDocumentError: "
                    f"{str(exc)[:120]}")


# ================================================================= DID NOT REPRODUCE (controls)
def test_control_del_nested_inside_ins_is_removed_with_its_author(tmp_path):
    """Round 5 attack that did NOT reproduce. Text inserted and then deleted — Word stores it
    as w:del inside w:ins — is correctly removed, and the revision author is blanked."""
    out, _ = redact(tmp_path, "nr_delins", [
        '<w:p><w:ins w:id="1" w:author="Advokat" w:date="2026-01-01T00:00:00Z">'
        '<w:r><w:t xml:space="preserve">Predavajuci </w:t></w:r>'
        '<w:del w:id="2" w:author="Koncipient" w:date="2026-01-02T00:00:00Z">'
        '<w:r><w:delText xml:space="preserve">Mária Kováčová</w:delText></w:r></w:del>'
        '<w:r><w:t xml:space="preserve"> suhlasi.</w:t></w:r></w:ins></w:p>'])
    assert surfaces_with(out, "Mária Kováčová") == []
    assert surfaces_with(out, "Koncipient") == []


def test_control_vml_textbox_with_no_drawingml_twin_is_redacted(tmp_path):
    """Did NOT reproduce. A bare legacy v:shape / v:textbox, the shape LibreOffice and Word 97
    converters emit, is reached by the paragraph walk and tagged `textbox`."""
    out, lm = redact(tmp_path, "nr_vml", [
        '<w:p><w:r><w:pict><v:shape id="tb1" type="#_x0000_t202" '
        'style="width:200pt;height:50pt"><v:textbox><w:txbxContent><w:p><w:r>'
        '<w:t xml:space="preserve">Kupujuci: Mária Kováčová</w:t></w:r></w:p>'
        '</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>'])
    assert surfaces_with(out, "Mária Kováčová") == []
    assert lm.occurrences["[MENO_1]"] == [("textbox", "Mária Kováčová")]


def test_control_smarttag_and_customxmlelement_split_a_name_cleanly(tmp_path):
    """Did NOT reproduce as a leak. A name split across a w:smartTag or a w:customXmlElement
    boundary WITHOUT splitting the run is still found and removed (round 4 predicted this for
    w:smartTag; w:customXmlElement is measured here for the first time). Only the
    w:smartTagPr/w:attr VALUE survives, and that is R5-07."""
    out, lm = redact(tmp_path, "nr_st", [
        '<w:p><w:r><w:t xml:space="preserve">Predavajuci Ján </w:t></w:r>'
        '<w:customXmlElement w:uri="urn:x" w:element="party">'
        '<w:r><w:t xml:space="preserve">Novák</w:t></w:r></w:customXmlElement>'
        '<w:r><w:t xml:space="preserve"> suhlasi.</w:t></w:r></w:p>'])
    assert surfaces_with(out, "Ján Novák") == []
    assert lm.occurrences["[MENO_1]"] == [("body", "Ján Novák")]


def test_control_comment_replies_are_redacted_and_both_authors_blanked(tmp_path):
    """Did NOT reproduce. A threaded comment (w15:commentEx parent/child) is an ordinary
    w:comment per reply: text redacted, w:author and w:initials blanked, and word/people.xml
    stripped of both the display name and the sign-in address."""
    comments = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:comments {NS}>'
        '<w:comment w:id="1" w:author="JUDr. Ján Novák" w:initials="JN" '
        'w:date="2026-01-01T00:00:00Z"><w:p w14:paraId="11111111"><w:r>'
        '<w:t xml:space="preserve">Overit RC klientky 855612/7788.</w:t></w:r></w:p></w:comment>'
        '<w:comment w:id="2" w:author="Mgr. Mária Kováčová" w:initials="MK" '
        'w:date="2026-01-02T00:00:00Z"><w:p w14:paraId="22222222"><w:r>'
        '<w:t xml:space="preserve">Klientka Mária Kováčová to potvrdila.</w:t></w:r></w:p>'
        '</w:comment></w:comments>')
    comments_ex = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w15:commentsEx {NS}><w15:commentEx w15:paraId="11111111" w15:done="0"/>'
        '<w15:commentEx w15:paraId="22222222" w15:paraIdParent="11111111" w15:done="0"/>'
        '</w15:commentsEx>')
    people = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w15:people {NS}>'
        '<w15:person w15:author="JUDr. Ján Novák"><w15:presenceInfo w15:providerId="AD" '
        'w15:userId="jan.novak@advokat.sk"/></w15:person>'
        '<w15:person w15:author="Mgr. Mária Kováčová"><w15:presenceInfo w15:providerId="AD" '
        'w15:userId="maria.kovacova@advokat.sk"/></w15:person></w15:people>')
    ms = "http://schemas.microsoft.com/office/2011/relationships"

    def patch(src, dst):
        return patch_package(
            src, dst,
            add={"word/comments.xml": comments.encode(),
                 "word/commentsExtended.xml": comments_ex.encode(),
                 "word/people.xml": people.encode()},
            content_types=[override("/word/comments.xml", "comments"),
                           override("/word/commentsExtended.xml", "commentsExtended"),
                           override("/word/people.xml", "people")],
            doc_rels=[relationship("rIdCm", "comments", "comments.xml"),
                      f'<Relationship Id="rIdCmx" Type="{ms}/commentsExtended" '
                      'Target="commentsExtended.xml"/>',
                      f'<Relationship Id="rIdPpl" Type="{ms}/people" Target="people.xml"/>'])

    out, lm = redact(tmp_path, "nr_cm", [
        '<w:p><w:commentRangeStart w:id="1"/>'
        '<w:r><w:t xml:space="preserve">Zmluvna strana.</w:t></w:r>'
        '<w:commentRangeEnd w:id="1"/><w:r><w:commentReference w:id="1"/></w:r></w:p>'],
        patch=patch)
    for needle in ("Ján Novák", "Mária Kováčová", "855612/7788", "jan.novak@advokat.sk"):
        assert surfaces_with(out, needle) == [], needle
    assert lm.occurrences["[RODNE_CISLO_1]"] == [("comment", "855612/7788")]


def test_control_a_hyperlink_relationship_inside_a_footnote_part_is_scrubbed(tmp_path):
    """Did NOT reproduce. _scrub_rel_targets sweeps EVERY .rels part in the saved package, so a
    w:hyperlink r:id in a footnote resolves through word/_rels/footnotes.xml.rels and is
    scrubbed there. (The FIELD-CODE spelling of the same link is R5-04.)"""
    footnotes = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:footnotes {NS}><w:footnote w:id="1"><w:p>'
        '<w:hyperlink r:id="rIdL"><w:r><w:t xml:space="preserve">kontakt</w:t></w:r>'
        '</w:hyperlink></w:p></w:footnote></w:footnotes>')
    fn_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f'<Relationship Id="rIdL" Type="{REL}/hyperlink" '
        'Target="mailto:maria.kovacova@gmail.com" TargetMode="External"/></Relationships>')

    def patch(src, dst):
        return patch_package(src, dst,
                             add={"word/footnotes.xml": footnotes.encode(),
                                  "word/_rels/footnotes.xml.rels": fn_rels.encode()},
                             content_types=[override("/word/footnotes.xml", "footnotes")],
                             doc_rels=[relationship("rIdFn", "footnotes", "footnotes.xml")])

    out, _ = redact(tmp_path, "nr_fnrel", [
        '<w:p><w:r><w:t>Telo.</w:t></w:r><w:r><w:footnoteReference w:id="1"/></w:r></w:p>'],
        patch=patch)
    assert surfaces_with(out, "maria.kovacova@gmail.com") == []
    assert "https://removed.invalid/" in part(out, "word/_rels/footnotes.xml.rels")


def test_control_nested_repeating_section_content_controls_are_redacted(tmp_path):
    """Did NOT reproduce. A w15:repeatingSection whose items are themselves w:sdt is reached by
    the descendant paragraph walk; each repeated party gets its own label."""
    out, lm = redact(tmp_path, "nr_rep", [
        '<w:sdt><w:sdtPr><w:id w:val="5"/><w:alias w:val="Zoznam stran"/>'
        '<w15:repeatingSection/></w:sdtPr><w:sdtContent>'
        '<w:sdt><w:sdtPr><w:id w:val="6"/><w15:repeatingSectionItem/></w:sdtPr><w:sdtContent>'
        '<w:p><w:r><w:t xml:space="preserve">Strana 1: Ján Novák</w:t></w:r></w:p>'
        '</w:sdtContent></w:sdt>'
        '<w:sdt><w:sdtPr><w:id w:val="7"/><w15:repeatingSectionItem/></w:sdtPr><w:sdtContent>'
        '<w:p><w:r><w:t xml:space="preserve">Strana 2: Mária Kováčová</w:t></w:r></w:p>'
        '</w:sdtContent></w:sdt></w:sdtContent></w:sdt>'])
    assert surfaces_with(out, "Ján Novák") == []
    assert surfaces_with(out, "Mária Kováčová") == []
    assert set(lm.occurrences) == {"[MENO_1]", "[MENO_2]"}


def test_control_a_main_part_at_a_non_standard_name_is_still_redacted(tmp_path):
    """Did NOT reproduce. OPC lets the main document part be named anything; python-docx
    resolves it through _rels/.rels, so the writer redacts it normally. (The extractor reports
    it as `other_xml_parts` rather than `document_xml` — a surface-naming nuance, not a leak.)"""
    src = doc_with_body(['<w:p><w:r><w:t>Predavajuci: Ján Novák.</w:t></w:r></w:p>'],
                        tmp_path / "nr_alt.docx")
    with zipfile.ZipFile(src) as z:
        items = {n: z.read(n) for n in z.namelist()}
    items["word/telo.xml"] = items.pop("word/document.xml")
    items["word/_rels/telo.xml.rels"] = items.pop("word/_rels/document.xml.rels")
    items["_rels/.rels"] = items["_rels/.rels"].replace(b"word/document.xml", b"word/telo.xml")
    items["[Content_Types].xml"] = items["[Content_Types].xml"].replace(
        b"/word/document.xml", b"/word/telo.xml")
    patched = str(tmp_path / "nr_alt_p.docx")
    with zipfile.ZipFile(patched, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in items.items():
            z.writestr(n, b)
    out = str(tmp_path / "nr_alt_r.docx")
    lm = redact_docx_collect(patched, out, list(KNOWN))
    assert surfaces_with(out, "Ján Novák") == []
    assert lm.occurrences["[MENO_1]"] == [("body", "Ján Novák")]
