"""v1.1 Phase F (CONTRACTS_v11.md §10): the report filename must carry the SOURCE format.

The v1 rule was ``<out_path stem>_report.txt``. A law office exporting ``zmluva.docx`` and
``zmluva.pdf`` in ONE batch therefore got two outputs (``zmluva_anon.docx``,
``zmluva_anon.pdf``) and ONE report file — the second export silently overwrote the first, so
the record of a document that had just been redacted was destroyed with no error anywhere.
That is a liability failure (context.md §9), not a cosmetic one, which is why it gets its own
file: the assertion that matters is that the FIRST report still contains its own content after
the second export has run.

RED at HEAD: ``report_path_for`` returns the extension-less name, so the two paths collide and
the surviving file holds the PDF's rows.

Hand-built fixtures (python-docx / fitz), no corpus import.
"""
from __future__ import annotations

import os

import docx as _docx
import fitz

from gui.model import export_file, out_path_for
from writer.decisions import RedactionDecisions
from writer.docx_body import redact_docx_body
from writer.pdf_body import redact_pdf
from writer.report import report_path_for

# Auto-redacted in both writers, and distinct per format so each report is identifiable.
_DOCX_TEXT = "Konatel Jan Novak, DIC 2023456789."
_PDF_TEXT = "Konatel Pavol Horak, DIC 2023456789."


def test_report_path_carries_the_source_extension() -> None:
    assert os.path.basename(report_path_for(r"C:\x\zmluva_anon.docx")) == "zmluva_anon_docx_report.txt"
    assert os.path.basename(report_path_for(r"C:\x\zmluva_anon.pdf")) == "zmluva_anon_pdf_report.txt"
    # Same directory as the output, always.
    assert os.path.dirname(report_path_for(r"C:\x\zmluva_anon.pdf")) == r"C:\x"
    # A docx and a pdf sharing a stem no longer collide — the whole point.
    assert report_path_for(r"C:\x\z_anon.docx") != report_path_for(r"C:\x\z_anon.pdf")


def test_same_stem_docx_and_pdf_write_two_distinct_reports(tmp_path) -> None:
    """The writers themselves: one stem, two formats, two surviving reports."""
    src_docx = tmp_path / "zmluva.docx"
    d = _docx.Document()
    d.add_paragraph(_DOCX_TEXT)
    d.save(str(src_docx))

    src_pdf = tmp_path / "zmluva.pdf"
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), _PDF_TEXT, fontname="helv", fontsize=11)
    doc.save(str(src_pdf))
    doc.close()

    out_docx = str(tmp_path / "zmluva_anon.docx")
    out_pdf = str(tmp_path / "zmluva_anon.pdf")

    redact_docx_body(str(src_docx), out_docx, known_entities=["Jan Novak"])
    rep_docx = report_path_for(out_docx)
    docx_report_before = open(rep_docx, encoding="utf-8").read()

    # The colliding export: same stem, other format, run SECOND.
    redact_pdf(str(src_pdf), out_pdf, known_entities=["Pavol Horak"])
    rep_pdf = report_path_for(out_pdf)

    assert rep_docx != rep_pdf
    assert os.path.exists(rep_docx) and os.path.exists(rep_pdf)
    # Neither overwrote the other: the docx report is byte-unchanged and each names its own
    # locations ("body" for docx, "page_1" for pdf).
    assert open(rep_docx, encoding="utf-8").read() == docx_report_before
    assert "body" in docx_report_before
    assert "page_1" in open(rep_pdf, encoding="utf-8").read()
    assert "page_1" not in docx_report_before


def test_export_file_uses_report_path_for_not_its_own_rule(tmp_path) -> None:
    """gui.model.export_file must RETURN the writer's own report path — it used to rebuild the
    rule by hand, which is how the two would drift apart again."""
    src = tmp_path / "zmluva.docx"
    d = _docx.Document()
    d.add_paragraph(_DOCX_TEXT)
    d.save(str(src))

    out, report = export_file(str(src), ["Jan Novak"], RedactionDecisions())
    assert out == out_path_for(str(src))
    assert report == report_path_for(out)
    assert os.path.basename(report) == "zmluva_anon_docx_report.txt"
    assert os.path.exists(report), "the returned path must be the file the writer actually wrote"
