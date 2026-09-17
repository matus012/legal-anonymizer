"""R4-X3: a file the writers cannot open is REFUSED by name, not by library traceback.

``context.md`` §3 promises that a document the tool cannot process is refused with a message
the reviewer can act on. That promise held for a PDF with no text layer, and since red-team
round 3 for one whose text cannot be decoded or is shredded. It did not hold for a file that
cannot be OPENED at all.

Measured before this landed:

  * a ``.docx`` with a malformed ``.rels`` part  -> ``lxml.etree.XMLSyntaxError``
  * a ``.docx`` that is not a ZIP                -> ``PackageNotFoundError``
  * a password-protected PDF                     -> ``ValueError: document closed or encrypted``
  * a truncated PDF                              -> ``FileDataError``

The application never crashed on these — ``gui/worker.py`` isolates each file and the review
screen shows the exception text. That IS the problem: it shows a message written for a
developer, which does not say what to do and does not answer the question the lawyer actually
has, **whether a half-redacted file is now on disk**.

Both refusals therefore state outright that nothing was written, and both are raised before
any output is created — which is what makes that statement true rather than merely reassuring.
"""
import os
import tempfile
import zipfile

import fitz
import pytest
from docx import Document

from writer.docx_body import redact_docx_body
from writer.errors import PasswordProtectedError, UnreadableDocumentError
from writer.pdf_body import redact_pdf


def _good_docx(path):
    doc = Document()
    doc.add_paragraph("Predávajúci: Ján Novák, rodné číslo 850315/0018")
    doc.save(path)


def _corrupt_rels(src, dst):
    with zipfile.ZipFile(src) as zin:
        items = [(n, zin.read(n)) for n in zin.namelist()]
    with zipfile.ZipFile(dst, "w") as zout:
        for name, data in items:
            zout.writestr(name, b"<Relationships><broken" if name.endswith(".rels") else data)


def test_a_docx_with_a_malformed_rels_part_is_refused_by_name(tmp_path):
    good = tmp_path / "good.docx"
    bad = tmp_path / "bad.docx"
    _good_docx(str(good))
    _corrupt_rels(str(good), str(bad))
    with pytest.raises(UnreadableDocumentError):
        redact_docx_body(str(bad), str(tmp_path / "out.docx"))


def test_a_docx_that_is_not_a_zip_is_refused_by_name(tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"this is not a docx at all")
    with pytest.raises(UnreadableDocumentError):
        redact_docx_body(str(bad), str(tmp_path / "out.docx"))


def test_nothing_is_written_when_a_docx_is_refused(tmp_path):
    """The claim the message makes, asserted rather than trusted."""
    bad = tmp_path / "bad.docx"
    out = tmp_path / "out.docx"
    bad.write_bytes(b"not a docx")
    with pytest.raises(UnreadableDocumentError):
        redact_docx_body(str(bad), str(out))
    assert not out.exists()


def test_a_truncated_pdf_is_refused_by_name(tmp_path):
    src = tmp_path / "t.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(str(src))
    doc.close()
    data = src.read_bytes()
    src.write_bytes(data[: len(data) // 2])
    with pytest.raises(UnreadableDocumentError):
        redact_pdf(str(src), str(tmp_path / "out.pdf"))


def test_a_password_protected_pdf_gets_its_OWN_refusal(tmp_path):
    """Named separately because the remedy is entirely within the lawyer's reach — open it,
    enter the password, save an unprotected copy. The Slovak land registry issues
    password-protected PDFs, so this is ordinary, not exotic."""
    src = tmp_path / "enc.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Predavajuci: Jan Novak", fontname="helv", fontsize=11)
    doc.save(str(src), encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="user")
    doc.close()
    with pytest.raises(PasswordProtectedError):
        redact_pdf(str(src), str(tmp_path / "out.pdf"))


def test_the_password_refusal_is_a_subclass():
    """So every caller that already refuses an unopenable document refuses this one too."""
    assert issubclass(PasswordProtectedError, UnreadableDocumentError)


def test_the_original_library_error_is_kept_for_the_bug_report(tmp_path):
    """The human sentence goes to the lawyer; the library's own words are what a maintainer
    needs, so they are carried on the exception rather than thrown away."""
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a docx")
    with pytest.raises(UnreadableDocumentError) as excinfo:
        redact_docx_body(str(bad), str(tmp_path / "out.docx"))
    assert excinfo.value.detail
    assert excinfo.value.__cause__ is not None


def test_an_ordinary_document_is_unaffected(tmp_path):
    """The quiet half."""
    src = tmp_path / "ok.docx"
    out = tmp_path / "out.docx"
    _good_docx(str(src))
    redact_docx_body(str(src), str(out), known_entities=["Ján Novák"])
    assert out.exists()
    assert "850315/0018" not in "\n".join(p.text for p in Document(str(out)).paragraphs)


def test_the_gui_messages_say_no_file_was_written():
    from gui.model import (
        MSG_PASSWORD_PROTECTED,
        MSG_RESIDUAL_CONTENT,
        MSG_UNREADABLE_DOC,
    )

    msgs = (MSG_UNREADABLE_DOC, MSG_PASSWORD_PROTECTED, MSG_RESIDUAL_CONTENT)
    assert len(set(msgs)) == len(msgs), "each refusal needs its OWN message, not a shared one"
    for msg in msgs:
        assert "NEBOL vytvorený" in msg, "the message must answer 'is there a file on disk?'"


def test_the_residual_content_refusal_reaches_the_lawyer_in_slovak(tmp_path, monkeypatch):
    """The clause fires, and the reviewer sees the Slovak message rather than the exception.

    ResidualContentError shipped with NO handler in gui/model.py, so it fell through to
    gui/worker.py's per-file ``except Exception`` and the review screen showed the exception's
    English sentence. That is precisely the failure writer/errors.py exists to prevent -- its
    docstring calls a library message "written for a developer" the problem, and this one was
    written by us.

    This refusal is the hardest of the set to explain, because the PAGE LOOKS RIGHT: the text is
    gone where the reader can see it and survives where only a tool can. So the message has to
    say where it would have been, and that it is recoverable.
    """
    import gui.model as model
    from writer.errors import ResidualContentError

    src = tmp_path / "residue.pdf"
    src.write_bytes(b"%PDF-1.7\n")

    # Patch the collect step itself, NOT the writer functions it calls. My first version of this
    # test patched names that do not exist on the module, so the real code ran, failed to open a
    # nine-byte file and returned MSG_UNREADABLE_DOC -- a green-looking red herring that would
    # have proved nothing about this clause.
    def _raise(*a, **k):
        raise ResidualContentError(str(src), ["Maria Kovacova"], [7])

    monkeypatch.setattr(model, "_collect", _raise)

    scan = model.scan_file(str(src), [])
    assert scan.error == model.MSG_RESIDUAL_CONTENT, (
        f"the lawyer sees {scan.error!r}, not the Slovak refusal"
    )
    assert scan.rows == []
